"""Excel admission workflow: all database/Graph calls are mocked."""

from io import BytesIO
import json
from urllib.parse import urlparse, parse_qs
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright

from verify_direct_admission import PREFIX, SUBJECTS, CATALOGS
from verify_runtime_optimizations import BASE_URL, ROOT, mock_context


SIMPLE_HEADERS = ["Tipo de documento *", "N\u00famero de c\u00e9dula *", "Primer nombre *",
                  "Segundo nombre", "Primer apellido *", "Segundo apellido"]


def simple_workbook(rows=None):
    workbook = Workbook()
    students = workbook.active
    students.title = "Estudiantes"
    students.append(SIMPLE_HEADERS)
    for row in rows or []:
        students.append(row)
    parameters = workbook.create_sheet("Par\u00e1metros")
    parameters.append(["Formato", "INTEC_INGRESO_DIRECTO_2"])
    parameters.append(["C\u00f3digo de carrera", "7"])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def verify(browser, name, viewport):
    context, page, _, _, errors = mock_context(browser, ["matricula-acad/ingreso-directo"], "matricula-acad", viewport)
    state = {"fail": True, "hold": False}
    held, saves, analyses, downloads = [], [], [], []
    items = []
    for number in range(2):
        student = {"identificacion": f"AB12345{number}", "tipo_documento": 2, "apellidos": "PEREZ LOPEZ",
                   "nombres": f"ANA {number}", "correo": f"ana{number}@example.test", "sexo": 2, "estado_civil": 1, "etnia": 1}
        items.append({"fila": number + 2, "identificacion": student["identificacion"],
                      "nombre_estudiante": student["apellidos"] + " " + student["nombres"], "correo": student["correo"],
                      "errores": [], "payload": {"solicitud_id": str(uuid4()), "estudiante": student, "matricula": {}}})
    items.append({"fila": 4, "identificacion": "invalid", "nombre_estudiante": "CORREO DUPLICADO",
                  "correo": "dup@example.test", "errores": ["Correo repetido en las filas 4, 5."], "payload": None})

    def successful(payload):
        student = payload["estudiante"]
        return {"ok": True, "solicitud_id": payload["solicitud_id"], "codigo_estud": "99",
                "identificacion": student["identificacion"], "nombre_estudiante": student["apellidos"] + " " + student["nombres"],
                "nivel": 1, "message": "Ingreso registrado", "documentos_pendientes": True, "matricula": {"inserted": 1}}

    def handler(route):
        path = urlparse(route.request.url).path
        query = parse_qs(urlparse(route.request.url).query)
        status = 200
        if path == PREFIX + "/catalog":
            body = {"datos_catalogos": CATALOGS, "carreras": [
                {"cod_anio_basica": "7", "nombre_basica": "Ciberseguridad"},
                {"cod_anio_basica": "8", "nombre_basica": "Otra carrera"}],
                "periodos": [{"codigo_periodo": "1060", "detalle_periodo": "C1-2026-PB", "tipo_matricula": "R"}],
                "jornadas": [{"value": "1", "label": "Matutina"}]}
        elif path == PREFIX + "/pensum":
            body = {"items": SUBJECTS if query["cod_anio_basica"] == ["7"] else []}
        elif path == PREFIX + "/excel/plantilla":
            downloads.append(query)
            route.fulfill(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", body=simple_workbook())
            return
        elif path == PREFIX + "/excel/validar":
            content = route.request.post_data
            assert 'name="file"' in content and '"cod_anio_basica":7' in content and '"materia_codes":[11]' in content
            analyses.append(content)
            for item in items[:2]:
                item["payload"]["matricula"] = {"cod_anio_basica": 7, "codigo_periodo": 1060, "nivel": 1,
                                               "materia_codes": [11], "cod_jornada": 1, "paralelo": "A"}
            body = {"carrera": {"cod_anio_basica": "7", "nombre_basica": "Ciberseguridad"},
                    "total": 3, "validos": 2, "invalidos": 1, "items": items}
        elif path == PREFIX + "/save":
            assert route.request.headers.get("content-type") == "application/json"
            payload = route.request.post_data_json
            saves.append(payload)
            assert payload["matricula"]["cod_anio_basica"] == 7 and payload["matricula"]["materia_codes"] == [11]
            if state["fail"]:
                status = 500
                body = {"detail": "Error temporal simulado"}
            elif state["hold"]:
                held.append((route, payload))
                return
            else:
                body = successful(payload)
        elif path.startswith(PREFIX) and path.endswith("/credenciales"):
            assert route.request.method == "GET"
            body = {"estado_general": "PENDIENTE", "datos_credenciales": {"nombres": "ANA", "apellidos": "PEREZ LOPEZ"}}
        elif path == "/api/document-expedients/context":
            assert query["identification"] == ["AB123450"]
            body = {"student": {"code": 99, "identification": "AB123450", "name": "PEREZ LOPEZ ANA 0", "career": "Ciberseguridad"},
                    "total_documents": 0, "total_expedients": 0, "expedients": []}
        else:
            raise AssertionError(path)
        route.fulfill(status=status, content_type="application/json", body=json.dumps(body))

    context.route("**/api/students/ingreso-directo/**", handler)
    context.route("**/api/document-expedients/**", handler)
    page.goto(BASE_URL)
    page.get_by_role("heading", name="ingreso_intec", exact=True).wait_for()
    page.get_by_role("button", name="Ingreso desde Excel", exact=True).click()
    area = page.locator(".direct-admission-excel")
    form = page.locator(".direct-admission form")
    expect(area.get_by_role("button", name="Descargar plantilla", exact=True)).to_be_disabled()
    form.get_by_label("Carrera", exact=False).select_option("7")
    expect(form.get_by_label("Nivel", exact=False)).to_have_value("1")
    with page.expect_download() as download:
        area.get_by_role("button", name="Descargar plantilla", exact=True).click()
    assert download.value.suggested_filename == "ingreso_intec_carrera_7.xlsx"
    with open(download.value.path(), "rb") as downloaded:
        workbook = load_workbook(downloaded)
        assert workbook["Par\u00e1metros"]["B2"].value == "7"
        assert workbook["Estudiantes"].max_column == 6
        assert [cell.value for cell in workbook["Estudiantes"][1]] == SIMPLE_HEADERS
        workbook.close()
    form.get_by_label("Per\u00edodo", exact=False).select_option("1060")
    area.get_by_label("Estudiantes (.xlsx)", exact=True).set_input_files({"name": "students.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": b"mocked"})
    expect(area.get_by_role("button", name="Validar Excel", exact=True)).to_be_disabled()
    form.get_by_role("checkbox", name="Matricular Matem\u00e1tica aplicada", exact=True).check()
    area.get_by_role("button", name="Validar Excel", exact=True).click()
    expect(area.get_by_role("checkbox", name="Ingresar fila 4", exact=True)).to_be_disabled()
    expect(area.get_by_role("button", name="Ingresar seleccionados", exact=True)).to_be_disabled()
    area.get_by_role("checkbox", name="Ingresar fila 2", exact=True).check()
    assert not saves
    page.screenshot(path=str(ROOT / ".runlogs" / f"direct-admission-excel-{name}-rows.png"), full_page=True)
    area.get_by_role("button", name="Ingresar seleccionados", exact=True).click()
    dialog = page.get_by_role("dialog", name="Confirmar ingreso desde Excel", exact=True)
    expect(dialog.locator("li")).to_have_count(1)
    assert dialog.evaluate("element => element.matches(':modal')")
    assert "C1-2026-PB" in dialog.inner_text()
    page.screenshot(path=str(ROOT / ".runlogs" / f"direct-admission-excel-{name}-confirm.png"))
    dialog.get_by_role("button", name="Confirmar ingresos", exact=True).evaluate("element => { element.click(); element.click(); }")
    area.get_by_text("Error temporal simulado", exact=True).wait_for()
    assert len(saves) == 1 and saves[0]["estudiante"]["identificacion"] == "AB123450"

    # Retry selected row with exact UUID/body; pause before advancing to the next student.
    state["fail"] = False
    state["hold"] = True
    area.get_by_role("button", name="Seleccionar v\u00e1lidos", exact=True).click()
    area.get_by_role("button", name="Reintentar seleccionados", exact=True).click()
    dialog.get_by_role("button", name="Confirmar ingresos", exact=True).click()
    page.wait_for_timeout(100)
    assert held and len(saves) == 2 and saves[0] == saves[1]
    expect(form.get_by_label("Carrera", exact=False)).to_be_disabled()
    expect(page.get_by_role("button", name="Ingresos registrados", exact=True)).to_be_disabled()
    area.get_by_role("button", name="Pausar despu\u00e9s del estudiante actual", exact=True).click()
    route, payload = held.pop()
    state["hold"] = False
    route.fulfill(content_type="application/json", body=json.dumps(successful(payload)))
    area.get_by_text("Proceso pausado · 1 de 2", exact=True).wait_for()
    assert len(saves) == 2
    expect(area.get_by_role("checkbox", name="Ingresar fila 2", exact=True)).to_be_disabled()
    area.get_by_role("button", name="Continuar seleccionados", exact=True).click()
    expect(dialog.locator("li")).to_have_count(1)
    dialog.get_by_role("button", name="Confirmar ingresos", exact=True).click()
    area.get_by_text("Proceso finalizado · 1 de 1", exact=True).wait_for()
    expect(area.locator(".direct-admission-excel-ok")).to_have_count(2)
    assert len(saves) == 3 and saves[-1]["estudiante"]["identificacion"] == "AB123451"
    page.screenshot(path=str(ROOT / ".runlogs" / f"direct-admission-excel-{name}-complete.png"), full_page=True)

    # Career changes invalidate the old analysis. Restoring the same draft keeps results.
    form.get_by_label("Carrera", exact=False).select_option("8")
    expect(area.get_by_role("checkbox", name="Ingresar fila 2", exact=True)).to_have_count(0)
    expect(area.get_by_role("button", name="Validar Excel", exact=True)).to_be_disabled()
    form.get_by_label("Carrera", exact=False).select_option("7")
    expect(form.get_by_label("Nivel", exact=False)).to_have_value("1")
    form.get_by_role("checkbox", name="Matricular Matem\u00e1tica aplicada", exact=True).check()
    expect(area.locator(".direct-admission-excel-ok")).to_have_count(2)
    area.get_by_role("button", name="Documentos", exact=True).first.click()
    page.get_by_role("heading", name="Documentaci\u00f3n del estudiante", exact=True).wait_for()
    page.get_by_role("button", name="Ingreso desde Excel", exact=True).click()
    expect(area.locator(".direct-admission-excel-ok")).to_have_count(2)
    assert len(saves) == 3 and len(analyses) == 1 and len(downloads) == 1
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth + 1")
    assert not errors, errors
    context.close()
    print(f"{name}: career-bound download, row validation/selection, idempotent retry, pause/resume and documents passed")


if __name__ == "__main__":
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="msedge", headless=True)
        for name, viewport in [("desktop", {"width": 1600, "height": 1000}),
                               ("mobile", {"width": 390, "height": 844}),
                               ("mobile-small", {"width": 320, "height": 740})]:
            verify(browser, name, viewport)
        browser.close()
