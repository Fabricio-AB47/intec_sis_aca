"""Office/Moodle provisioning UI; academic, Graph and Moodle requests are mocked."""

from io import BytesIO
from email.parser import BytesParser
from email.policy import default
import json
from urllib.parse import urlparse
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright

from verify_direct_admission import CATALOGS, PREFIX, SUBJECTS
from verify_direct_admission_excel import SIMPLE_HEADERS, simple_workbook
from verify_runtime_optimizations import BASE_URL, ROOT, mock_context


NAMES = {"primer_nombre": "ANA MARIA", "segundo_nombre": "", "primer_apellido": "DE LA CRUZ", "segundo_apellido": "PEREZ"}
NAME = "DE LA CRUZ PEREZ ANA MARIA"
EMAIL = "anamaria.delacruz@intec.edu.ec"


def verify(browser, name, viewport):
    context, page, _, _, errors = mock_context(browser, ["matricula-acad/ingreso-directo"], "matricula-acad", viewport)
    saves, provisions = [], []
    records, states, attempts = {}, {}, {}
    receipts = []
    receipt_failure = {"single": True, "bulk": True}
    batch_id = str(uuid4())

    def result(status):
        return {"estado_general": status, "estado_graph": "CREADO_GRAPH", "estado_licencia": "ASIGNADA_ESTUDIANTE",
                "estado_moodle": "ERROR_MOODLE" if status == "PARCIAL" else "EXISTENTE_MOODLE", "correo_institucional": EMAIL,
                "datos_persona": NAMES, "reporte_credencial_id": 88, "estado_correo_academico": "SINCRONIZADO",
                "errores": ["Moodle temporalmente no disponible; la matr\u00edcula sigue guardada."] if status == "PARCIAL" else [],
                "observacion": "Las contrase\u00f1as de cuentas existentes no fueron modificadas."}

    def handler(route):
        path = urlparse(route.request.url).path
        status = 200
        if path == PREFIX + "/catalog":
            body = {"datos_catalogos": CATALOGS, "carreras": [{"cod_anio_basica": "7", "nombre_basica": "Ciberseguridad"}],
                    "periodos": [{"codigo_periodo": "1060", "detalle_periodo": "C1-2026-PB", "tipo_matricula": "R"}],
                    "jornadas": [{"value": "1", "label": "Matutina"}]}
        elif path == PREFIX + "/pensum":
            body = {"items": SUBJECTS}
        elif path == PREFIX + "/preview":
            assert route.request.post_data_json["credenciales"]["primer_apellido"].upper() == "DE LA CRUZ"
            body = {"summary": {"insertar": 1}, "items": [{**SUBJECTS[0], "accion": "INSERTAR"}],
                    "estudiante": {"accion": "EXISTENTE", "codigo_estud": "99"}}
        elif path == PREFIX + "/save":
            payload = route.request.post_data_json
            saves.append(payload)
            assert payload["credenciales"]["primer_nombre"].upper() == "ANA MARIA"
            if payload["solicitud_id"] == batch_id:
                assert payload["estudiante"]["correo"] == ""
                assert (payload["estudiante"]["sexo"], payload["estudiante"]["estado_civil"], payload["estudiante"]["etnia"]) == (3, 6, 9)
            body = {"ok": True, "solicitud_id": payload["solicitud_id"], "codigo_estud": str(98 + len(saves)),
                    "identificacion": payload["estudiante"]["identificacion"], "nombre_estudiante": NAME, "nivel": 1,
                    "message": "Matr\u00edcula registrada correctamente", "matricula": {"inserted": 1}, "credenciales": NAMES,
                    "datos_credenciales": {"nombres": "ANA MARIA", "apellidos": "DE LA CRUZ PEREZ"}, "estudiante_existente": True}
            records[payload["solicitud_id"]] = body
        elif path.endswith("/comprobante") or path == PREFIX + "/comprobantes":
            bulk = path == PREFIX + "/comprobantes"
            ids = route.request.post_data_json["solicitud_ids"] if bulk else [path.split("/")[-2]]
            assert all(identifier in records for identifier in ids)
            assert ids == [batch_id] if bulk else ids == [saves[0]["solicitud_id"]]
            kind = "bulk" if bulk else "single"
            if receipt_failure[kind]:
                route.fulfill(status=503, content_type="application/json", body=json.dumps({"detail": "PDF temporalmente no disponible; la matr\u00edcula sigue guardada."}))
            else:
                receipts.append(ids)
                route.fulfill(content_type="application/pdf", headers={"Cache-Control": "no-store, private"}, body=b"%PDF-1.4\n%%EOF")
            return
        elif path.endswith("/credenciales"):
            request_id = path.split("/")[-2]
            if route.request.method == "GET":
                body = states.get(request_id, {"estado_general": "PENDIENTE", "datos_persona": NAMES})
            else:
                assert route.request.post_data_json == NAMES
                provisions.append((request_id, route.request.post_data_json))
                attempts[request_id] = attempts.get(request_id, 0) + 1
                if request_id == batch_id and attempts[request_id] == 1:
                    status = 503
                    body = {"detail": "Respuesta temporal de Office 365; reintente credenciales sin repetir matr\u00edcula."}
                else:
                    body = result("PARCIAL" if request_id != batch_id and attempts[request_id] == 1 else "COMPLETO")
                    states[request_id] = body
        elif path == PREFIX + "/history":
            body = {"items": [{**record, "carrera": "Ciberseguridad", "periodo": "C1-2026-PB", "registrado_por": "admin@example.test"} for record in records.values()]}
        elif path == PREFIX + "/excel/validar":
            message = BytesParser(policy=default).parsebytes(
                b"Content-Type: " + route.request.headers["content-type"].encode() + b"\r\nMIME-Version: 1.0\r\n\r\n" + route.request.post_data_buffer)
            parts = {part.get_param("name", header="content-disposition"): part for part in message.iter_parts()}
            assert parts["crear_credenciales"].get_payload(decode=True) == b"true"
            enrollment = json.loads(parts["matricula"].get_payload(decode=True))
            assert enrollment["cod_anio_basica"] == 7 and enrollment["materia_codes"] == [11]
            attachment = parts["file"]
            workbook = load_workbook(BytesIO(attachment.get_payload(decode=True)))
            assert [cell.value for cell in workbook["Estudiantes"][1]] == SIMPLE_HEADERS
            assert [cell.value for cell in workbook["Estudiantes"][2]] == [2, "CD123456", "ANA MARIA", None, "DE LA CRUZ", "PEREZ"]
            workbook.close()
            body = {"carrera": {"cod_anio_basica": "7", "nombre_basica": "Ciberseguridad"}, "total": 1, "validos": 1, "invalidos": 0,
                    "items": [{"fila": 2, "identificacion": "CD123456", "nombre_estudiante": NAME, "correo": "", "errores": [],
                               "estudiante_existente": True, "codigo_estud_existente": "100",
                               "payload": {"solicitud_id": batch_id, "estudiante": {"identificacion": "CD123456", "tipo_documento": 2, "nombres": "ANA MARIA", "apellidos": "DE LA CRUZ PEREZ", "correo": "", "sexo": 3, "estado_civil": 6, "etnia": 9},
                                           "matricula": {"cod_anio_basica": 7, "codigo_periodo": 1060, "nivel": 1, "materia_codes": [11], "paralelo": "A", "cod_jornada": 1}, "credenciales": NAMES}}]}
        elif path == "/api/document-expedients/context":
            body = {"student": {"code": 99, "identification": "1724036536", "name": NAME}, "expedients": [], "total_documents": 0, "total_expedients": 0}
        elif path == "/api/admin/credenciales/history/88/report":
            workbook = Workbook(); workbook.active.append(["Correo", "Documento de prueba"]); workbook.active.append([EMAIL, "mock"])
            buffer = BytesIO(); workbook.save(buffer)
            route.fulfill(content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", body=buffer.getvalue()); return
        else:
            raise AssertionError(path)
        route.fulfill(status=status, content_type="application/json", body=json.dumps(body))

    context.route("**/api/students/ingreso-directo/**", handler)
    context.route("**/api/document-expedients/**", handler)
    context.route("**/api/admin/credenciales/**", handler)
    page.goto(BASE_URL)
    page.get_by_role("heading", name="ingreso_intec", exact=True).wait_for()
    modes = page.get_by_role("navigation", name="Modalidad de matriculaci\u00f3n", exact=True)
    expect(modes.get_by_role("button", name="ingreso_intec", exact=False)).to_have_count(1)
    expect(modes.get_by_role("button", name="ingreso_intec", exact=False)).to_have_attribute("aria-current", "page")
    if viewport["width"] < 1000:
        page.get_by_role("button", name="Abrir men\u00fa principal", exact=True).click()
    menu = page.get_by_role("navigation", name="Men\u00fa principal", exact=True)
    section = menu.locator(".student-nav__section").filter(has=page.get_by_text("Matr\u00edcula", exact=True))
    expect(section).to_have_count(1)
    if section.locator(".student-nav__group-button").get_attribute("aria-expanded") != "true":
        section.locator(".student-nav__group-button").click()
    entry = section.get_by_role("button", name="ingreso_intec", exact=False)
    expect(entry).to_have_count(1)
    expect(menu.get_by_role("button", name="Ingreso directo", exact=False)).to_have_count(0)
    page.screenshot(path=str(ROOT / ".runlogs" / f"ingreso-intec-{name}-menu.png"))
    entry.click()
    expect(page.get_by_role("heading", name="ingreso_intec", exact=True)).to_be_visible()
    form = page.locator(".direct-admission form")
    form.get_by_label("Identificaci\u00f3n", exact=False).click()
    for label, value in [("Identificaci\u00f3n", "1724036536"), ("Nombres", "ANA MARIA"), ("Apellidos", "DE LA CRUZ PEREZ"), ("Correo personal", "ana@example.test")]:
        form.get_by_label(label, exact=False).fill(value)
    for label, value in [("Sexo", "2"), ("Estado civil", "1"), ("Etnia", "1"), ("Carrera", "7"), ("Per\u00edodo", "1060")]:
        form.get_by_label(label, exact=False).select_option(value)
    form.get_by_role("checkbox", name="Crear / verificar Office 365 y Moodle despu\u00e9s de matricular", exact=True).check()
    for label, value in [("Primer nombre Office 365", "ANA MARIA"), ("Segundo nombre Office 365", ""), ("Primer apellido Office 365", "DE LA CRUZ"), ("Segundo apellido Office 365", "PEREZ")]:
        form.get_by_label(label, exact=False).fill(value)
    form.get_by_role("checkbox", name="Matricular Matem\u00e1tica aplicada", exact=True).check()
    form.get_by_role("button", name="Validar matr\u00edcula", exact=True).click()
    dialog = page.get_by_role("dialog", name="Confirmar ingreso y matr\u00edcula", exact=True)
    assert "DE LA CRUZ" in dialog.inner_text()
    expect(dialog.get_by_text("Estudiante existente · Código 99", exact=True)).to_be_visible()
    dialog.get_by_role("button", name="Registrar y matricular", exact=True).click()
    credentials_area = page.get_by_role("region", name="Credenciales del estudiante", exact=True)
    credentials_area.get_by_role("status").filter(has_text="PARCIAL").wait_for()
    assert len(saves) == len(provisions) == 1
    page.get_by_role("status").filter(has_text="Matrícula registrada correctamente").wait_for()
    page.screenshot(path=str(ROOT / ".runlogs" / f"direct-credentials-{name}-partial.png"), full_page=True)
    credentials_area.get_by_role("button", name="Crear / verificar credenciales", exact=True).evaluate("element => { element.click(); element.click(); }")
    credentials_area.get_by_role("status").filter(has_text="COMPLETO").wait_for()
    assert len(saves) == 1 and len(provisions) == 2
    assert page.locator('.direct-admission').get_by_text('clave_permanente', exact=False).count() == 0
    credentials_area.get_by_role("button", name="Verificar cuentas", exact=True).evaluate("element => { element.click(); element.click(); }")
    expect(credentials_area.get_by_role("button", name="Verificar cuentas", exact=True)).to_be_enabled()
    assert len(provisions) == 3 and len(saves) == 1
    expect(credentials_area.get_by_text("SINCRONIZADO", exact=True)).to_be_visible()
    credentials_area.get_by_role("button", name="Descargar comprobante PDF", exact=True).click()
    expect(credentials_area.get_by_role("alert")).to_contain_text("PDF temporalmente no disponible")
    receipt_failure["single"] = False
    with page.expect_download() as receipt:
        credentials_area.get_by_role("button", name="Descargar comprobante PDF", exact=True).click()
    assert receipt.value.suggested_filename == "comprobante_matricula_1724036536.pdf"
    assert len(provisions) == 3 and len(saves) == 1
    with page.expect_download() as download:
        credentials_area.get_by_role("button", name="Descargar credenciales", exact=True).click()
    assert download.value.suggested_filename == "credenciales_1724036536.xlsx"
    page.get_by_role("button", name="Ingresos registrados", exact=True).click()
    page.get_by_role("button", name="Documentos", exact=True).click()
    credentials_area.get_by_role("status").filter(has_text="COMPLETO").wait_for()
    assert len(provisions) == 3

    # Excel retries provision only; already saved academic enrollment is not sent again.
    page.get_by_role("button", name="Ingreso desde Excel", exact=True).click()
    area = page.locator(".direct-admission-excel")
    area.get_by_role("checkbox", name="Crear / verificar Office 365 y Moodle por cada estudiante ingresado", exact=True).check()
    area.get_by_label("Estudiantes (.xlsx)", exact=True).set_input_files({"name": "students.xlsx", "mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "buffer": simple_workbook([[2, "CD123456", "ANA MARIA", "", "DE LA CRUZ", "PEREZ"]])})
    area.get_by_role("button", name="Validar Excel", exact=True).click()
    expect(area.get_by_text("Estudiante existente · Código 100", exact=True)).to_be_visible()
    expect(area.get_by_role("button", name="Descargar comprobantes PDF", exact=True)).to_be_disabled()
    area.get_by_role("checkbox", name="Ingresar fila 2", exact=True).check()
    area.get_by_role("button", name="Ingresar seleccionados", exact=True).click()
    confirm = page.get_by_role("dialog", name="Confirmar ingreso desde Excel", exact=True)
    confirm.get_by_role("button", name="Confirmar ingresos", exact=True).click()
    area.get_by_text("Respuesta temporal de Office 365; reintente credenciales sin repetir matrícula.", exact=True).wait_for()
    assert len(saves) == 2 and len(provisions) == 4
    area.get_by_role("button", name="Reintentar seleccionados", exact=True).click()
    confirm.get_by_role("button", name="Confirmar ingresos", exact=True).click()
    area.get_by_text("COMPLETO", exact=True).wait_for()
    expect(area.get_by_role("checkbox", name="Ingresar fila 2", exact=True)).to_be_disabled()
    assert len(saves) == 2 and len(provisions) == 5
    expect(area.get_by_text("Existente verificado · Código 100", exact=True)).to_be_visible()
    assert provisions[-1] == provisions[-2]
    area.get_by_role("button", name="Descargar comprobantes PDF", exact=True).click()
    expect(area.get_by_role("alert")).to_contain_text("PDF temporalmente no disponible")
    receipt_failure["bulk"] = False
    with page.expect_download() as receipt:
        area.get_by_role("button", name="Descargar comprobantes PDF", exact=True).click()
    assert receipt.value.suggested_filename == "comprobantes_matricula.pdf"
    assert len(saves) == 2 and len(provisions) == 5 and len(receipts) == 2
    assert page.locator('.direct-admission').get_by_text('clave_permanente', exact=False).count() == 0
    page.screenshot(path=str(ROOT / ".runlogs" / f"direct-credentials-{name}-excel.png"), full_page=True)
    assert page.evaluate("document.documentElement.scrollWidth <= innerWidth + 1")
    assert not errors, errors
    context.close()
    print(f"{name}: automatic provisioning, compound names, partial retry, credential download and enrollment isolation passed")


if __name__ == "__main__":
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="msedge", headless=True)
        for name, viewport in [("desktop", {"width": 1600, "height": 1000}), ("mobile", {"width": 390, "height": 844}), ("mobile-small", {"width": 320, "height": 740})]:
            verify(browser, name, viewport)
        browser.close()
