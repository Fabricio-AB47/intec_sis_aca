"""SENESCYT teacher-model download with all API responses mocked."""

from io import BytesIO
import json
from urllib.parse import parse_qs, urlparse
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from playwright.sync_api import expect, sync_playwright

from verify_runtime_optimizations import BASE_URL, ROOT, mock_context


def workbook_content():
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active.append(["tipoDocumentoId", "numeroIdentificacion", "primerApellido"])
    workbook.active.append([1, "0401105002", "GUERRON"])
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def student_zip_content():
    workbook = Workbook()
    workbook.active.title = "Sheet1"
    workbook.active.append(["tipoDocumentoId", "numeroIdentificacion", "primerApellido"])
    workbook.active.append([1, "0106889843", "GALARZA"])
    xlsx = BytesIO()
    workbook.save(xlsx)
    workbook.close()
    archive_content = BytesIO()
    with ZipFile(archive_content, "w", ZIP_DEFLATED) as archive:
        archive.writestr("01_Administracion_Financiera_completo.xlsx", xlsx.getvalue())
    return archive_content.getvalue()


def teacher_zip_content():
    archive_content = BytesIO()
    with ZipFile(archive_content, "w", ZIP_DEFLATED) as archive:
        archive.writestr("01_Administracion_completo.xlsx", workbook_content())
    return archive_content.getvalue()


def verify(browser, name, viewport):
    context, page, _, _, errors = mock_context(browser, ["senescyt-estudiantes"], "senescyt-estudiantes", viewport)
    requests = []
    queries = []
    exports = []

    def handler(route):
        parsed = urlparse(route.request.url)
        params = parse_qs(parsed.query)
        if parsed.path == "/api/students/senescyt/catalogo":
            route.fulfill(content_type="application/json", body=json.dumps({
                "careers": [], "targets": ["estudiantes", "docentes"], "export_modes": ["completo", "faltantes"],
                "periods": [
                    {"codigo_periodo": 1032, "nombre_periodo": "C2-HOMO-2025-PB", "fecha_inicio": "2025-10-01"},
                    {"codigo_periodo": 1033, "nombre_periodo": "C2-2025-PC OCTUBRE 2025 - MARZO 2026", "fecha_inicio": "2025-10-01"},
                    {"codigo_periodo": 1060, "nombre_periodo": "C1-2026-PCFF", "fecha_inicio": "2026-04-01"},
                ],
            }))
        elif parsed.path == "/api/students/senescyt/datos":
            queries.append(params)
            route.fulfill(content_type="application/json", body=json.dumps({"target": params.get("target", ["estudiantes"])[0], "summary": {}, "rows": [], "careers": [], "missing_fields": []}))
        elif parsed.path == "/api/students/senescyt/datos/export":
            target = params["target"][0]
            mode = params["mode"][0]
            requests.append((target, mode))
            exports.append(params)
            if (target, mode) == ("docentes", "completo"):
                route.fulfill(content_type="application/zip", body=teacher_zip_content())
            elif (target, mode) == ("estudiantes", "completo"):
                route.fulfill(content_type="application/zip", body=student_zip_content())
            else:
                route.fulfill(content_type="application/zip", body=b"mock-zip")
        else:
            raise AssertionError(parsed.path)

    context.route("**/api/students/senescyt/**", handler)
    page.goto(BASE_URL)
    page.get_by_role("heading", name="Datos SENESCYT", exact=True).wait_for()
    target_select = page.locator(".senescyt-target-control select")
    target_select.select_option("docentes", timeout=10000)
    complete = page.get_by_role("button", name="Archivo docentes por carrera", exact=True)
    expect(complete).to_be_visible()
    expect(page.get_by_role("heading", name="Generación por carrera y faltantes", exact=True)).to_be_visible()
    page.screenshot(path=str(ROOT / ".runlogs" / f"senescyt-teacher-model-{name}.png"), full_page=True)
    with page.expect_download() as download:
        complete.click()
    assert download.value.suggested_filename == "senescyt-docentes-completo-todas-las-carreras.zip"
    with ZipFile(download.value.path()) as archive:
        assert archive.namelist() == ["01_Administracion_completo.xlsx"]
        workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])), read_only=True)
        assert workbook.sheetnames == ["Sheet1"]
        workbook.close()
    with page.expect_download() as missing:
        page.get_by_role("button", name="Faltantes docentes global/carreras", exact=True).click()
    assert missing.value.suggested_filename.endswith(".zip")
    target_select.select_option("estudiantes")
    expect(page.get_by_text("Cada ZIP contiene un Excel por carrera con el modelo SENESCYT de estudiantes.", exact=False)).to_be_visible()
    with page.expect_download() as student:
        page.get_by_role("button", name="Archivo estudiantes por carrera", exact=True).click()
    assert student.value.suggested_filename.endswith(".zip")
    with ZipFile(student.value.path()) as archive:
        assert archive.namelist() == ["01_Administracion_Financiera_completo.xlsx"]
        workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])), read_only=True)
        assert workbook.sheetnames == ["Sheet1"]
        assert workbook.active["B2"].value == "0106889843"
        workbook.close()
    assert requests == [("docentes", "completo"), ("docentes", "faltantes"), ("estudiantes", "completo")]
    assert all("periodo" not in query and "fecha_limite" not in query for query in exports)

    page.locator(".senescyt-period-picker summary").click()
    page.get_by_label("Buscar período", exact=True).fill("2025")
    picker = page.locator(".senescyt-period-picker")
    picker.get_by_role("button", name="Seleccionar visibles", exact=True).click()
    expect(picker.get_by_role("checkbox", checked=True)).to_have_count(2)
    expect(picker.locator("summary")).to_contain_text("2 período(s)")
    page.get_by_label("Fecha límite (inclusive)", exact=False).fill("2025-12-31")
    expect(page.get_by_role("button", name="Archivo estudiantes por carrera", exact=True)).to_be_disabled()
    expect(page.get_by_role("button", name="Vista previa", exact=True)).to_be_disabled()
    page.get_by_role("button", name="Consultar", exact=True).click()
    expect(page.get_by_role("button", name="Archivo estudiantes por carrera", exact=True)).to_be_enabled()
    page.get_by_role("button", name="Vista previa", exact=True).click()
    expect(page.get_by_role("dialog")).to_be_visible()
    page.get_by_role("button", name="Cerrar", exact=True).click()
    for target in ("estudiantes", "docentes"):
        target_select.select_option(target)
        for label in (f"Archivo {target} por carrera", f"Faltantes {target} global/carreras"):
            with page.expect_download():
                page.get_by_role("button", name=label, exact=True).click()
            assert exports[-1]["periodo"] == ["1032", "1033"]
            assert exports[-1]["fecha_limite"] == ["2025-12-31"]
        assert queries[-1]["periodo"] == ["1032", "1033"]
        assert queries[-1]["fecha_limite"] == ["2025-12-31"]
    page.screenshot(path=str(ROOT / ".runlogs" / f"senescyt-period-filters-{name}.png"), full_page=True)
    assert page.evaluate("document.documentElement.scrollWidth <= window.innerWidth"), "Horizontal overflow"
    picker.get_by_role("button", name="Limpiar", exact=True).click()
    page.get_by_label("Fecha límite (inclusive)", exact=False).fill("")
    expect(picker.locator("summary")).to_contain_text("Todos los períodos")
    page.get_by_role("button", name="Consultar", exact=True).click()
    expect(page.get_by_role("button", name="Archivo docentes por carrera", exact=True)).to_be_enabled()
    assert "periodo" not in queries[-1] and "fecha_limite" not in queries[-1]
    assert not errors, errors
    context.close()
    print(f"{name}: model ZIPs, periods, inclusive cutoff, pending filters and reset passed")


if __name__ == "__main__":
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(channel="msedge", headless=True)
        for name, viewport in [("desktop", {"width": 1500, "height": 940}), ("mobile", {"width": 390, "height": 844})]:
            verify(browser, name, viewport)
        browser.close()
