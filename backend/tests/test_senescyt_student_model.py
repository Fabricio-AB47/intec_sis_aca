from datetime import date, datetime
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pandas as pd

from app.core.security import SessionUser
from app.routers import senescyt


HEADERS = """tipoDocumentoId numeroIdentificacion primerApellido segundoApellido primerNombre segundoNombre sexoId generoId estadocivilId etniaId pueblonacionalidadId tipoSangre discapacidad porcentajeDiscapacidad numCarnetConadis tipoDiscapacidad fechaNacimiento paisNacionalidadId provinciaNacimientoId cantonNacimientoId paisResidenciaId provinciaResidenciaId cantonResidenciaId tipoColegioId modalidadCarrera jornadaCarrera fechaInicioCarrera fechaMatricula tipoMatriculaId nivelAcademicoQueCursa duracionPeriodoAcademico haRepetidoAlMenosUnaMateria paraleloId haPerdidoLaGratuidad recibePensionDiferenciada estudianteocupacionId ingresoEstudianteId bonoDesarrolloId haRealizadoPracticasPreprofesionales nroHorasPracticasPreprofesionalesPorPeriodo entornoInstitucionalPracticasProfesionales sectorEconomicoPracticaProfesional tipoBecaId primeraRazonBecaId segundaRazonBecaId terceraRazonBecaId cuartaRazonBecaId quintaRazonBecaId sextaRazonBecaId montoBeca porcientoBecaCoberturaArancel porcientoBecaCoberturaManuntencion financiamientoBeca montoAyudaEconomica montoCreditoEducativo participaEnProyectoVinculacionSociedad tipoAlcanceProyectoVinculacionId correoElectronico numeroCelular nivelFormacionPadre nivelFormacionMadre ingresoTotalHogar cantidadMiembrosHogar""".split()


def rows():
    return pd.DataFrame([
        {"codigo": "123", "nombreCarrera": "Administración Financiera", "nombreCompleto": "GALARZA CANDO TANIA FERNANDA",
         "tipoDocumentoId": 1, "numeroIdentificacion": "0106889843", "primerApellido": "GALARZA",
         "segundoApellido": "CANDO", "primerNombre": "TANIA", "segundoNombre": "FERNANDA",
         "fechaNacimiento": date(1995, 1, 12), "fechaInicioCarrera": date(2025, 9, 9),
         "fechaMatricula": datetime(2025, 9, 9, 13, 45), "provinciaNacimientoId": "01", "cantonNacimientoId": "0103",
         "ingresosestudianteId": 2, "bonodesarrolloId": 2, "correoElectronico": "tania@intec.edu.ec  ",
         "numeroCelular": "0963090648", "ingresoTotalHogar": 600, "cantidadMiembrosHogar": 2},
        {"codigo": "456", "nombreCarrera": "Administración", "primerApellido": "PEREZ",
         "primerNombre": "ANA", "numeroIdentificacion": "0805575735", "correoElectronico": "=1+1"},
    ])


def report(dataframe=None):
    return {"target": "estudiantes", "dataframe": rows() if dataframe is None else dataframe,
            "report_columns": senescyt._REPORT_COLUMNS}


def test_student_model_has_exact_supplied_headers_values_and_no_audit_columns():
    assert len(HEADERS) == 63
    assert senescyt._STUDENT_MODEL_COLUMNS == HEADERS
    content = senescyt._student_model_workbook(rows().head(1))
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == ["Sheet1"]
    sheet = workbook.active
    assert sheet.max_row == 2 and sheet.max_column == 63
    assert [cell.value for cell in sheet[1]] == HEADERS
    assert all(cell.font.bold and cell.alignment.horizontal == "center" and cell.border.left.style == "thin"
               and cell.border.right.style == "thin" for cell in sheet[1])
    values = dict(zip(HEADERS, (cell.value for cell in sheet[2])))
    assert values["numeroIdentificacion"] == "0106889843"
    assert values["provinciaNacimientoId"] == "01" and values["cantonNacimientoId"] == "0103"
    assert values["numeroCelular"] == "0963090648"
    assert values["fechaNacimiento"] == "1995-01-12"
    assert values["fechaInicioCarrera"] == values["fechaMatricula"] == "2025-09-09"
    assert values["correoElectronico"] == "tania@intec.edu.ec"
    assert values["ingresoEstudianteId"] == values["bonoDesarrolloId"] == 2
    assert values["pueblonacionalidadId"] is None
    assert "codigo" not in HEADERS and "nombreCarrera" not in HEADERS
    workbook.close()


def test_student_model_blocks_formula_execution_and_preserves_empty_template():
    workbook = load_workbook(BytesIO(senescyt._student_model_workbook(rows().tail(1))), data_only=False)
    email = workbook.active.cell(2, HEADERS.index("correoElectronico") + 1)
    assert email.value == "=1+1" and email.data_type == "s"
    workbook.close()
    workbook = load_workbook(BytesIO(senescyt._student_model_workbook(pd.DataFrame())))
    assert workbook.sheetnames == ["Sheet1"]
    assert [cell.value for cell in workbook.active[1]] == HEADERS
    assert workbook.active.max_row == 1
    workbook.close()


def test_student_model_preserves_leading_zeroes_when_sql_returns_numeric_locations():
    data = rows().head(1).copy()
    data.loc[data.index[0], "provinciaNacimientoId"] = 1
    data.loc[data.index[0], "cantonNacimientoId"] = 103.0
    data.loc[data.index[0], "provinciaResidenciaId"] = 7
    data.loc[data.index[0], "cantonResidenciaId"] = 702
    workbook = load_workbook(BytesIO(senescyt._student_model_workbook(data)))
    values = dict(zip(HEADERS, (cell.value for cell in workbook.active[2])))
    assert values["provinciaNacimientoId"] == "01"
    assert values["cantonNacimientoId"] == "0103"
    assert values["provinciaResidenciaId"] == "07"
    assert values["cantonResidenciaId"] == "0702"
    workbook.close()


def test_complete_student_zip_contains_only_model_workbooks_per_career():
    content = senescyt._audit_export_zip(report(), "completo")
    with ZipFile(BytesIO(content)) as archive:
        assert len(archive.namelist()) == 2
        for filename in archive.namelist():
            workbook = load_workbook(BytesIO(archive.read(filename)))
            assert workbook.sheetnames == ["Sheet1"]
            assert [cell.value for cell in workbook.active[1]] == HEADERS
            assert workbook.active.max_row == 2
            workbook.close()


def test_student_zip_keeps_latest_matriculation_once_per_career():
    data = rows().head(1).copy()
    previous = data.copy()
    previous["fechaMatricula"] = date(2024, 5, 1)
    content = senescyt._audit_export_zip(report(pd.concat([previous, data], ignore_index=True)), "completo")
    with ZipFile(BytesIO(content)) as archive:
        workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])))
        assert workbook.active.max_row == 2
        assert workbook.active.cell(2, HEADERS.index("fechaMatricula") + 1).value == "2025-09-09"
        workbook.close()


def test_student_guide_context_uses_na_only_when_an_answer_supports_it():
    data = rows().head(1).copy()
    data["discapacidad"] = 2
    data["haRealizadoPracticasPreprofesionales"] = 2
    data["tipoBecaId"] = 3
    data["participaEnProyectoVinculacionSociedad"] = 2
    data["paisNacionalidadId"] = "57"
    data["ingresoTotalHogar"] = 0
    data["montoAyudaEconomica"] = 0
    workbook = load_workbook(BytesIO(senescyt._student_model_workbook(data)))
    values = dict(zip(HEADERS, (cell.value for cell in workbook.active[2])))
    assert values["porcentajeDiscapacidad"] == "NA"
    assert values["numCarnetConadis"] == "NA"
    assert values["tipoDiscapacidad"] == 7
    assert values["nroHorasPracticasPreprofesionalesPorPeriodo"] == "NA"
    assert values["entornoInstitucionalPracticasProfesionales"] == 5
    assert values["sectorEconomicoPracticaProfesional"] == 22
    assert values["montoBeca"] == "NA"
    assert values["primeraRazonBecaId"] == 2
    assert values["financiamientoBeca"] == 4
    assert values["tipoAlcanceProyectoVinculacionId"] == 5
    assert values["provinciaNacimientoId"] == values["cantonNacimientoId"] == "NA"
    assert values["ingresoTotalHogar"] == values["montoAyudaEconomica"] == "NA"
    workbook.close()


def test_student_normalization_does_not_invent_codes_for_missing_source_values():
    raw = {column: None for column in senescyt._REPORT_COLUMNS}
    raw.update({"codigoEstud": "7", "nombreCarrera": "Administracion",
                "Apellidos_nombre": "PEREZ LOPEZ ANA MARIA"})
    normalized = senescyt._normalize_dataframe(pd.DataFrame([raw])).iloc[0]
    for column in ("tipoDocumentoId", "tipoSangre", "tipoColegioId", "bonodesarrolloId", "tipoBecaId",
                   "pueblonacionalidadId", "paraleloId", "sectorEconomicoPracticaProfesional"):
        assert pd.isna(normalized[column]), column
    assert pd.isna(normalized["ingresoTotalHogar"])
    assert senescyt._audit_field_filled(pd.Series({"fechaMatricula": "31/02/2025"}), "fechaMatricula", "estudiantes") is False
    assert senescyt._audit_field_filled(pd.Series({"numeroCelular": "098123456"}), "numeroCelular", "estudiantes") is False
    assert senescyt._audit_field_filled(pd.Series({"correoElectronico": "correo-invalido"}), "correoElectronico", "estudiantes") is False
    assert not senescyt._audit_field_filled(pd.Series({"tipoSangre": 9}), "tipoSangre", "estudiantes")
    assert senescyt._audit_field_filled(pd.Series({"tipoSangre": 7}), "tipoSangre", "estudiantes")


def test_student_disability_card_uses_the_seven_digit_guide_format():
    assert senescyt._normalize_conadis("0123456") == "0123456"
    assert senescyt._normalize_conadis("0123456789") == "NA"
    assert senescyt._audit_field_filled(
        pd.Series({"discapacidad": 1, "numCarnetConadis": "0123456"}), "numCarnetConadis", "estudiantes"
    )
    assert not senescyt._audit_field_filled(
        pd.Series({"discapacidad": 1, "numCarnetConadis": "0123456789"}), "numCarnetConadis", "estudiantes"
    )
    assert senescyt._document_analysis(pd.Series({"tipoDocumentoId": 2, "numeroIdentificacion": "ab1234567"}))["valido"]


def test_selected_career_is_exact_and_audit_rows_match_export_deduplication():
    data = rows()
    selected = senescyt._filter_by_career(data, ["Administración"])
    assert selected["nombreCarrera"].tolist() == ["Administración"]
    repeated = pd.DataFrame([
        {"codigo": "123", "numeroIdentificacion": "0106889843", "nombreCarrera": "Administración Financiera",
         "fechaMatricula": "2025-09-09"},
        {"codigo": "123", "numeroIdentificacion": "0106889843", "nombreCarrera": "Administración Financiera",
         "fechaMatricula": "2024-09-09"},
        {"codigo": "456", "numeroIdentificacion": "0805575735", "nombreCarrera": "Administración",
         "fechaMatricula": "2025-09-09"},
    ])
    deduplicated = senescyt._dedupe_career_model_rows(repeated, "fechaMatricula")
    assert len(deduplicated) == 2
    assert set(deduplicated["nombreCarrera"]) == {"Administración", "Administración Financiera"}


def test_empty_student_zip_keeps_model_headers():
    content = senescyt._audit_export_zip(report(pd.DataFrame()), "completo")
    with ZipFile(BytesIO(content)) as archive:
        assert len(archive.namelist()) == 1
        workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])))
        assert [cell.value for cell in workbook.active[1]] == HEADERS
        workbook.close()


def test_legacy_student_export_and_unified_export_share_the_same_model():
    app = FastAPI()
    app.include_router(senescyt.router)
    app.dependency_overrides[senescyt._SENESCYT_ACCESS] = lambda: SessionUser(login="admin@intec.edu.ec", rol="ADMINISTRADOR")
    legacy_report = {"dataframe": rows().rename(columns={"codigo": "codigoEstud"})}
    with TestClient(app) as client, \
         patch.object(senescyt, "_build_report", return_value=legacy_report), \
         patch.object(senescyt, "_build_senescyt_audit", return_value=report()):
        for url in ("/api/students/senescyt/estudiantes/export",
                    "/api/students/senescyt/datos/export?target=estudiantes&mode=completo"):
            response = client.get(url)
            assert response.status_code == 200 and response.headers["content-type"].startswith("application/zip")
            with ZipFile(BytesIO(response.content)) as archive:
                assert len(archive.namelist()) == 2
                for name in archive.namelist():
                    workbook = load_workbook(BytesIO(archive.read(name)))
                    assert workbook.sheetnames == ["Sheet1"]
                    assert [cell.value for cell in workbook.active[1]] == HEADERS
                    workbook.close()
