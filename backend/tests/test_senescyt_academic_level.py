from datetime import date
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pandas as pd
import pytest

from app.core.security import SessionUser
from app.routers import senescyt


def student_source(level):
    row = {field: None for field in senescyt._REPORT_COLUMNS if field not in senescyt._NAME_FIELDS}
    row.update({
        "codigoEstud": "123",
        "tipoDocumentoId": 1,
        "numeroIdentificacion": "0106889843",
        "Apellidos_nombre": "PEREZ LOPEZ ANA MARIA",
        "nombreCarrera": "Administracion",
        "fechaMatricula": date(2025, 10, 1),
        "nivelAcademicoQueCursa": level,
    })
    return pd.DataFrame([row])


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(senescyt.router)
    app.dependency_overrides[senescyt._SENESCYT_ACCESS] = lambda: SessionUser(
        login="admin@intec.edu.ec", rol="ADMINISTRADOR",
    )
    with TestClient(app) as instance:
        yield instance


def test_level_is_highest_numeric_semester_from_exact_career_and_enrolled_subjects():
    cutoff = date(2025, 12, 31)
    with patch.object(senescyt, "_read_sql_dataframe", return_value=pd.DataFrame()) as reader:
        senescyt._read_student_audit_source([1032, 1033], cutoff)
    sql, params = reader.call_args.args
    assert params == [1032, 1033, cutoff]
    assert "LEFT JOIN dbo.PENSUM pen ON pen.codigo_materia = cx.codigo_materia" in sql
    assert "AND pen.Cod_AnioBasica = cx.cod_anio_Basica" in sql
    assert "MAX(CASE WHEN TRY_CONVERT(int, pen.Semestre) > 0" in sql
    assert "THEN TRY_CONVERT(int, pen.Semestre) END) AS semestre" in sql
    assert "MAX(semestre) OVER (" in sql
    assert "PARTITION BY codigo_estud, cod_anio_Basica" in sql
    assert "cne.nivel_academico AS nivelAcademicoQueCursa" in sql
    assert "e.nivelAcademicoQueCursa" not in sql
    assert sql.index("p.cod_periodo IN (?, ?)") < sql.index("MAX(semestre) OVER")
    assert sql.index("cx.Fecha_Matricula) <= ?") < sql.index("MAX(semestre) OVER")
    assert sql.index("MAX(semestre) OVER") < sql.index("WHERE s.posicion = 1")
    assert "PromedioFinal" not in sql


@pytest.mark.parametrize("level", [1, 4, 9, None])
def test_preview_and_complete_or_missing_excel_use_same_computed_level(client, level):
    params = [("target", "estudiantes"), ("carrera", "Administracion"),
              ("periodo", 1032), ("periodo", 1033), ("fecha_limite", "2025-12-31")]
    with patch.object(senescyt, "_read_sql_dataframe", return_value=student_source(level)) as reader, patch.object(
        senescyt, "_read_dataframe", side_effect=AssertionError("Stored level must not be used in reports"),
    ):
        response = client.get("/api/students/senescyt/datos", params=params)
        assert response.status_code == 200, response.text
        row = response.json()["rows"][0]
        assert row["fields"]["nivelAcademicoQueCursa"] == level
        assert ("nivelAcademicoQueCursa" in row["campos_faltantes"]) == (level is None)

        for mode in ("completo", "faltantes"):
            response = client.get("/api/students/senescyt/datos/export", params=params + [("mode", mode)])
            assert response.status_code == 200, response.text
            with ZipFile(BytesIO(response.content)) as archive:
                workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])))
                matched = 0
                for sheet in workbook:
                    headers = [cell.value for cell in sheet[1]]
                    if "nivelAcademicoQueCursa" in headers:
                        column = headers.index("nivelAcademicoQueCursa") + 1
                        assert sheet.cell(2, column).value == level
                        matched += 1
                assert matched
                workbook.close()
        assert reader.call_count == 3
        assert all(call.args[1] == [1032, 1033, date(2025, 12, 31)] for call in reader.call_args_list)


def test_new_generation_recalculates_level_instead_of_reusing_previous_result(client):
    with patch.object(senescyt, "_read_sql_dataframe", side_effect=[student_source(2), student_source(4)]) as reader:
        first = client.get("/api/students/senescyt/datos?target=estudiantes&periodo=1032")
        second = client.get("/api/students/senescyt/datos?target=estudiantes&periodo=1032")
    assert first.status_code == second.status_code == 200
    assert first.json()["rows"][0]["fields"]["nivelAcademicoQueCursa"] == 2
    assert second.json()["rows"][0]["fields"]["nivelAcademicoQueCursa"] == 4
    assert reader.call_count == 2


def test_level_stays_separate_for_each_student_career(client):
    data = pd.concat([student_source(4), student_source(2).assign(nombreCarrera="Ciberseguridad")], ignore_index=True)
    with patch.object(senescyt, "_read_sql_dataframe", return_value=data):
        response = client.get("/api/students/senescyt/datos?target=estudiantes")
    assert response.status_code == 200
    levels = {row["nombre_carrera"]: row["fields"]["nivelAcademicoQueCursa"] for row in response.json()["rows"]}
    assert levels == {"Administracion": 4, "Ciberseguridad": 2}


def test_legacy_report_and_export_also_use_computed_level(client):
    with patch.object(senescyt, "_read_sql_dataframe", return_value=student_source(4)) as reader, patch.object(
        senescyt, "_count_scalar", return_value=1,
    ), patch.object(senescyt, "_read_dataframe", side_effect=AssertionError("Stored level must not be used")):
        report = client.get("/api/students/senescyt/estudiantes")
        assert report.status_code == 200, report.text
        response = client.get("/api/students/senescyt/estudiantes/export")
        assert response.status_code == 200, response.text
        with ZipFile(BytesIO(response.content)) as archive:
            workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])))
            headers = [cell.value for cell in workbook.active[1]]
            assert workbook.active.cell(2, headers.index("nivelAcademicoQueCursa") + 1).value == 4
            workbook.close()
        assert reader.call_count == 2
