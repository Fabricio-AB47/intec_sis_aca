import asyncio
from copy import deepcopy
from io import BytesIO
from unittest.mock import AsyncMock, MagicMock, patch

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from app.core.security import SessionUser
from app.integrations.moodle.exceptions import MoodleApiError
from app.routers import moodle as router
from app.services import moodle_enrollment_validation as validation
from app.services import moodle_academic_validation as academic_validation
from app.services.moodle_enrollment_validation_report import excel_report, pdf_report


def person(code=1, **values):
    return {"student_code": code, "student": f"ESTUDIANTE {code}", "document": f"001000000{code}",
            "state": "A", "email": f"estudiante{code}@intec.edu.ec", "registry_email": "", **values}


def user(code=1, **values):
    return {"id": code, "fullname": f"ESTUDIANTE {code}", "idnumber": f"001000000{code}",
            "email": f"ESTUDIANTE{code}@INTEC.EDU.EC", "username": "", "role_shortnames": ["student"], **values}


def subject(career=1, number=101, code="VGA-ES-2023-95"):
    return {"career_code": career, "career_name": f"CARRERA {career}", "subject_id": number,
            "subject_code": code, "subject_name": "Inteligencia Artificial"}


def enrollment(code=1, career=1, number=101, **values):
    return {"student_code": code, "career_code": career, "career": f"CARRERA {career}",
            "subject_id": number, "parallel": "PB1", "period_code": 10, "period": "PERIODO 10", **values}


def scope(users=None, **values):
    return {"course": {"id": 20, "fullname": "Inteligencia Artificial - PB1", "shortname": "VGA-ES-2023-95-R30"},
            "parallel": "PB1", "users": [user()] if users is None else users, "error": "", **values}


def report(**values):
    args = {"period": {"code": 10, "name": "PERIODO 10"}, "registry": [person()],
            "subjects": [subject()], "enrollments": [enrollment()], "history": [], "career_changes": [],
            "snapshots": [scope()], "actor": "admin@example.test", "warnings": []}
    return validation.compare_enrollments(**{**args, **values})


def statuses(result):
    return [r["status"] for r in result["rows"]]


def test_exact_match_and_previous_career_do_not_make_a_valid_enrollment_incorrect():
    result = report(history=[enrollment(career=2, period_code=9, period="PERIODO 9")])
    assert statuses(result) == ["COINCIDE"]
    assert result["rows"][0]["previous_careers"] == "CARRERA 2"
    assert result["rows"][0]["academic_career"] == "CARRERA 1"


def test_common_subject_in_multiple_careers_is_matched_using_enrollment_composite_key():
    result = report(subjects=[subject(), subject(2, 201)], enrollments=[enrollment(career=2, number=201)])
    assert statuses(result) == ["COINCIDE"]
    assert result["rows"][0]["academic_career"] == "CARRERA 2"


def test_finds_missing_students_in_both_directions():
    result = report(registry=[person(), person(2), person(3)], enrollments=[enrollment(), enrollment(2)],
                    snapshots=[scope([user(), user(3)])])
    assert set(statuses(result)) == {"COINCIDE", "FALTA_MOODLE", "SIN_MATRICULA_PERIODO"}
    assert next(r for r in result["rows"] if r["status"] == "FALTA_MOODLE")["student_code"] == "2"


def test_previous_career_subject_requires_review_not_current_career_reassignment():
    result = report(enrollments=[enrollment(career=2, number=201)], subjects=[subject(), subject(2, 201, "VGA-ES-2023-96")],
                    history=[enrollment(period_code=9, period="PERIODO 9")])
    assert statuses(result) == ["CAMBIO_CARRERA"]
    assert "PERIODO 9" in result["rows"][0]["reason"]


def test_applied_career_change_is_considered_when_old_enrollments_were_removed():
    result = report(enrollments=[enrollment(career=2, number=201)], career_changes=[
        {"student_code": 1, "origin_code": 1, "destination_code": 2, "origin": "CARRERA 1",
         "destination": "CARRERA 2", "applied_at": "2026-09-21"},
    ])
    assert statuses(result) == ["CAMBIO_CARRERA"]
    assert "2026-09-21" in result["rows"][0]["career_changes"]


def test_wrong_career_and_parallel_are_reported_without_fuzzy_name_matching():
    assert statuses(report(enrollments=[enrollment(career=2)])) == ["OTRA_CARRERA"]
    assert statuses(report(enrollments=[enrollment(parallel="PB2")])) == ["OTRO_PARALELO"]
    assert statuses(report(enrollments=[enrollment(parallel="PB2")], snapshots=[scope(parallel="*")])) == ["COINCIDE"]


@pytest.mark.parametrize("moodle_user, registry", [
    (user(username="otro@intec.edu.ec"), [person()]),
    (user(idnumber="0099999999"), [person()]),
    (user(), [person(), person(2, email="estudiante1@intec.edu.ec")]),
    (user(), [person(), person(2, document="0010000001")]),
    (user(), [person(), person(state="I")]),
])
def test_conflicting_or_duplicated_identity_never_creates_a_false_missing_result(moodle_user, registry):
    result = report(registry=registry, snapshots=[scope([moodle_user])])
    assert "IDENTIDAD_AMBIGUA" in statuses(result)
    assert "FALTA_MOODLE" not in statuses(result)


def test_missing_institutional_email_can_match_exact_passport_including_letters_and_zeroes():
    result = report(registry=[person(document="P00123", email="")],
                    snapshots=[scope([user(idnumber="P00123", email="externo@example.test")])])
    assert statuses(result) == ["COINCIDE"]
    assert statuses(report(registry=[person(document="P00123", email="")],
                           snapshots=[scope([user(idnumber="Q00123", email="externo@example.test")])])) == ["NO_IDENTIFICADO", "NO_VERIFICABLE"]


def test_moodle_errors_and_unknown_roles_do_not_prove_absence():
    assert set(statuses(report(snapshots=[scope([], error="No autorizado")]))) == {"ERROR_MOODLE", "NO_VERIFICABLE"}
    assert "FALTA_MOODLE" not in statuses(report(snapshots=[scope([user(role_shortnames=[])])]))


def test_inactive_students_are_not_expected_but_are_visible_when_still_in_moodle():
    assert statuses(report(registry=[person(state="I")])) == ["INACTIVO"]
    assert statuses(report(registry=[person(state="I")], snapshots=[scope([])])) == []
    assert statuses(report(snapshots=[scope([user(suspended=True)])])) == ["SUSPENDIDO_MOODLE"]


def test_multiple_accounts_and_duplicate_enrollments_are_not_accepted_as_matches():
    assert "DUPLICADO_MOODLE" in statuses(report(snapshots=[scope([user(), user(id=2)])]))
    assert "MATRICULA_DUPLICADA" in statuses(report(enrollments=[enrollment(), enrollment()]))
    assert statuses(report(enrollments=[enrollment(), enrollment()], snapshots=[scope([])])) == ["MATRICULA_DUPLICADA"]


def test_enrollment_header_uses_student_career_and_period_not_subject_attempt_number():
    assert statuses(report(enrollments=[enrollment(header_exists=0)])) == ["SIN_CABECERA"]
    assert statuses(report(enrollments=[enrollment(header_exists=0)], snapshots=[scope([])])) == ["SIN_CABECERA"]
    cursor = MagicMock()
    cursor.description = []
    cursor.fetchall.return_value = []
    validation.MoodleEnrollmentValidationService._enrollments(cursor, "cx.codigo_periodo = ?", [10])
    sql, value = cursor.execute.call_args.args
    assert "cab.codigo_estud = cx.codigo_estud" in sql
    assert "cab.cod_anio_Basica = cx.cod_anio_Basica" in sql
    assert "cab.codigo_periodo = cx.codigo_periodo" in sql
    assert "cab.Num_Matricula" not in sql
    assert value == 10


def test_multiple_aulas_are_checked_together_and_do_not_report_false_missing_students():
    first = scope([])
    second = deepcopy(scope())
    second["course"]["id"] = 21
    assert statuses(report(snapshots=[first, second])) == ["COINCIDE"]


def test_unmatched_and_conflicting_course_codes_are_reviewed():
    item = scope()
    item["course"]["shortname"] = "SIN-CODIGO"
    assert statuses(report(snapshots=[item])) == ["SIN_PENSUM"]
    item = scope()
    item["course"]["idnumber"] = "VGA-ES-2023-96-R30"
    result = report(subjects=[subject(), subject(2, 201, "VGA-ES-2023-96")], snapshots=[item])
    assert "SIN_PENSUM" in statuses(result)
    assert "COINCIDE" not in statuses(result)


@pytest.mark.parametrize("academic_code, moodle_code, method", [
    ("VGA-ES-2023-95", "VGA-ES-2023-95", "Exacta"),
    ("VGA-ES-2023-95", " vga / es _ 2023 - 095 - H12026 ", "Normalizada"),
    ("VGA-ES-2023-95-", "VGAES202395R30", "Normalizada"),
    ("VGAES202395", "VGA-ES-2023-95-H12026", "Normalizada"),
    ("vga es 2023 095", "VGA-ES-2023-95-R30", "Normalizada"),
    ("VGA-ES-2023-95", "VGA\u2013ES\u20132023\u201395\u2013R30", "Normalizada"),
    ("VGA-ES-2023-95", "\uff36\uff27\uff21-ES-2023-95-R30", "Normalizada"),
    ("VGA-ES-2023-95", "VGA-E5-2023-95-R30", "Similitud controlada"),
    ("VGA-ES-2023-95", "VGA-SE-2023-95-R30", "Similitud controlada"),
    ("VGA-ES-2023-95", "VWGA-DSDSWES-2023-95-R232026", "Similitud controlada"),
    ("VWGA-DSDSWES-2023-95", "VGA-ES-2023-95-R30", "Similitud controlada"),
])
def test_subject_codes_reconcile_formatting_and_controlled_noise_in_both_systems(academic_code, moodle_code, method):
    item = scope()
    item["course"]["shortname"] = moodle_code
    result = report(subjects=[subject(code=academic_code)], snapshots=[item])
    assert statuses(result) == ["COINCIDE"]
    assert result["rows"][0]["code_match"] == method
    assert result["rows"][0]["moodle_code"] == moodle_code.strip()
    assert result["rows"][0]["subject_code"] == academic_code
    assert result["scope"][0]["code_match_detail"]


def test_normalized_pensum_variants_preserve_the_students_actual_career_and_subject_code():
    result = report(subjects=[subject(), subject(2, 201, "VGA ES 2023 095-")],
                    enrollments=[enrollment(career=2, number=201)])
    assert statuses(result) == ["COINCIDE"]
    assert result["rows"][0]["subject_code"] == "VGA ES 2023 095-"
    assert result["rows"][0]["academic_career"] == "CARRERA 2"


def test_exact_code_takes_precedence_over_similar_codes_without_enrollment_bias():
    result = report(subjects=[subject(), subject(2, 201, "VGA-E5-2023-95")],
                    enrollments=[enrollment(career=2, number=201)])
    assert "COINCIDE" not in statuses(result)
    assert result["scope"][0]["subject_code"] == "VGA-ES-2023-95"


@pytest.mark.parametrize("code", ["VGA-ES-2023-96-R30", "VGA-ES-2023-59-R30", "VGA-ES-2023-9-R30", "VGA-ES-2025-95-R30"])
def test_numeric_code_typos_are_review_candidates_not_automatic_matches(code):
    item = scope()
    item["course"]["shortname"] = code
    result = report(snapshots=[item])
    assert set(statuses(result)) == {"SIN_PENSUM", "NO_VERIFICABLE"}
    assert "VGA-ES-2023-95" in result["rows"][0]["code_candidates"]
    assert result["rows"][0]["code_match"] == "Requiere revisión"


def test_multiple_similar_prefixes_are_never_resolved_by_score_or_student_enrollment():
    item = scope()
    item["course"]["shortname"] = "VGA-EZ-2023-95-R30"
    result = report(subjects=[subject(), subject(2, 201, "VGA-EX-2023-95")], snapshots=[item])
    assert set(statuses(result)) == {"SIN_PENSUM", "NO_VERIFICABLE"}
    assert "VGA-ES-2023-95" in result["rows"][0]["code_candidates"]
    assert "VGA-EX-2023-95" in result["rows"][0]["code_candidates"]


def test_identifier_conflicts_and_unrecognized_structured_id_are_visible_not_silently_ignored():
    for identifier in ["VGA-ES-2023-96-R30", "VGA-ES-2023-678-R30"]:
        item = scope()
        item["course"]["idnumber"] = identifier
        result = report(snapshots=[item])
        assert "COINCIDE" not in statuses(result)
        assert "FALTA_MOODLE" not in statuses(result)
        assert result["rows"][0]["moodle_idnumber"] == identifier


def test_identifiers_can_share_one_normalized_code_but_names_cannot_replace_subject_code():
    item = scope()
    item["course"]["idnumber"] = "VGAES2023095R30"
    assert statuses(report(snapshots=[item])) == ["COINCIDE"]
    item["course"]["shortname"] = "SIN-CODIGO"
    item["course"]["idnumber"] = ""
    assert statuses(report(snapshots=[item])) == ["SIN_PENSUM"]


def test_similar_code_reports_include_originals_and_reconciliation_evidence():
    item = scope()
    item["course"]["shortname"] = "VGA-E5-2023-95-R30"
    item["course"]["idnumber"] = "VGA-ES-2023-95-R30"
    result = report(snapshots=[item])
    assert statuses(result) == ["COINCIDE"]
    sheet = load_workbook(BytesIO(excel_report(result)))["Detalle"]
    cells = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
    assert cells["Comparación del código"] == "Similitud controlada"
    assert cells["Número ID curso Moodle"] == item["course"]["idnumber"]
    assert cells["Código Moodle"] == item["course"]["shortname"]
    content = "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf_report(result))).pages)
    assert "Similitud controlada" in content
    assert item["course"]["shortname"] in content and item["course"]["idnumber"] in content


def test_excel_is_text_safe_and_pdf_preserves_rows_and_scope():
    result = report(registry=[person(student="=FORMULA()")])
    book = load_workbook(BytesIO(excel_report(result)))
    assert book["Detalle"]["C2"].value == "=FORMULA()"
    assert book["Detalle"]["C2"].data_type == "s"
    assert book["Detalle"]["D2"].value == "0010000001"
    content = "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf_report(result))).pages)
    assert "0010000001" in content and "PERIODO 10" in content and "Coincide" in content


def test_pdf_handles_long_details_and_multiple_pages():
    result = report()
    result["rows"][0]["reason"] = "Observacion de prueba " * 700
    result["rows"] *= 4
    pages = PdfReader(BytesIO(pdf_report(result))).pages
    assert len(pages) > 1


def test_scan_is_sequential_read_only_and_exports_are_owner_bound_and_expire():
    moodle = MagicMock()
    moodle.get_all_courses = AsyncMock(return_value=[scope()["course"], {**scope()["course"], "id": 21}])
    moodle.get_course_enrolled_users = AsyncMock(side_effect=[MoodleApiError("private upstream detail"), [user()]])
    service = validation.MoodleEnrollmentValidationService(moodle)
    academic = {"period": {"code": 10, "name": "PERIODO 10"}, "subjects": [subject()],
                "registry": [person()], "enrollments": [enrollment()]}
    with patch.object(service, "_load_academic", return_value=academic), patch.object(service, "_load_history", return_value=([], [], [])):
        result = asyncio.run(service.validate(period_code=10, courses=[{"id": 20, "parallel": "PB1"}, {"id": 21, "parallel": "PB1"}], actor="admin"))
    assert [c.args[0] for c in moodle.get_course_enrolled_users.await_args_list] == [20, 21]
    assert all(c.kwargs["refresh"] for c in moodle.get_course_enrolled_users.await_args_list)
    assert "private upstream detail" not in str(result)
    assert service.report(result["report_id"], "admin") == result
    with pytest.raises(HTTPException):
        service.report(result["report_id"], "other")
    with patch.object(validation, "monotonic", return_value=validation.monotonic() + 1801):
        with pytest.raises(HTTPException):
            service.report(result["report_id"], "admin")


def test_routes_validate_scope_and_require_own_access():
    app = FastAPI()
    app.include_router(router.router)
    service = MagicMock()
    service.validate = AsyncMock(return_value={"rows": []})
    app.dependency_overrides[router.get_moodle_enrollment_validation_service] = lambda: service
    app.dependency_overrides[router._MOODLE_ENROLLMENT_VALIDATION_ACCESS] = lambda: SessionUser(login="admin", rol="ADMINISTRADOR")
    with TestClient(app) as client:
        valid = {"period_code": 10, "courses": [{"id": 20, "parallel": " pb1 "}]}
        assert client.post("/api/moodle/enrollment-validation/preview", json=valid).status_code == 200
        assert service.validate.await_args.kwargs["courses"] == [{"id": 20, "parallel": "PB1"}]
        for payload in [{**valid, "period_code": 0}, {**valid, "courses": []},
                        {**valid, "courses": valid["courses"] * 2},
                        {**valid, "courses": [{"id": 20, "parallel": " "}]}]:
            assert client.post("/api/moodle/enrollment-validation/preview", json=payload).status_code == 422
        def denied():
            raise HTTPException(403, "No autorizado")
        app.dependency_overrides[router._MOODLE_ENROLLMENT_VALIDATION_ACCESS] = denied
        assert client.post("/api/moodle/enrollment-validation/preview", json=valid).status_code == 403
        assert client.get("/api/moodle/enrollment-validation/reports/anything/pdf").status_code == 403


def academic_student(**overrides):
    subjects = overrides.pop("subjects", [subject()])
    people = validation._identity_indexes([person()])[0]
    return academic_validation.student_audit(**{
        "person": people[1], "career": {"code": 1, "name": "CARRERA 1"}, "period": {"code": 10, "name": "PERIODO 10"},
        "subjects": subjects, "enrollments": [enrollment()], "history": [], "changes": [],
        "snapshots": [{"course": scope()["course"], "error": ""}], "user": user(),
        "resolver": academic_validation.EnrollmentSubjectResolver(subjects), **overrides,
    })


def test_academic_view_relates_current_career_period_and_all_returned_courses_separately():
    subjects = [subject(), subject(2, 201, "VGA-ES-2023-96"), subject(3, 301, "VGA-ES-2023-97"), subject(1, 102, "VGA-ES-2023-98")]
    snapshots = [{"course": {**scope()["course"], "id": n, "shortname": f"VGA-ES-2023-{n}-R30"}, "error": ""} for n in [95, 96, 97, 98]]
    student = academic_student(subjects=subjects, snapshots=snapshots,
                               history=[enrollment(career=2, number=201, period_code=9, period="PERIODO ANTERIOR")])
    assert {r["status"] for r in student["moodle_courses"]} == {"MATERIA_MATRICULADA", "ANTECEDENTE_ACADEMICO", "OTRA_CARRERA", "SIN_MATRICULA_PERIODO"}
    assert student["academic_career"] == "CARRERA 1" and student["period_code"] == "10"
    assert student["academic_subjects"][0]["linked"]
    historic = next(r for r in student["rows"] if r["status"] == "ANTECEDENTE_ACADEMICO")
    assert "PERIODO ANTERIOR" in historic["enrollment_periods"]
    assert "CARRERA 2" in historic["subject_careers"]
    assert historic["period_relation"] == "Período del aula Moodle no confirmado"


def test_academic_view_checks_real_subject_enrollment_not_just_career_membership():
    student = academic_student(enrollments=[enrollment(number=102)], subjects=[subject(), subject(number=102, code="VGA-ES-2023-96")])
    assert student["moodle_courses"][0]["status"] == "SIN_MATRICULA_PERIODO"
    assert any(r["status"] == "SIN_AULA_ACTIVA" and r["subject_code"] == "VGA-ES-2023-96" for r in student["rows"])
    student = academic_student(subjects=[subject(), subject(2, 201)], enrollments=[enrollment(career=2, number=201)])
    assert student["rows"][0]["status"] == "OTRA_CARRERA_PERIODO"


def test_academic_view_includes_every_semester_and_unopened_modular_subjects():
    student = academic_student(subjects=[{**subject(), "level": 2}, {**subject(number=102, code="VGA-ES-2023-96"), "level": 1}],
                               enrollments=[enrollment(), enrollment(number=102)], snapshots=[])
    assert [s["semester"] for s in student["academic_subjects"]] == [1, 2]
    assert {r["status"] for r in student["rows"]} == {"SIN_AULA_ACTIVA"}
    assert all("apertura modular" in r["reason"] for r in student["rows"])


def test_academic_view_errors_and_ambiguous_codes_cannot_confirm_absence():
    student = academic_student(snapshots=[], problem="NO_VERIFICABLE", problem_detail="Error de consulta")
    assert {r["status"] for r in student["rows"]} == {"NO_VERIFICABLE"}
    student = academic_student(snapshots=[{"course": scope()["course"], "error": "Rol no confirmado"}])
    assert {r["status"] for r in student["rows"]} == {"NO_VERIFICABLE"}
    item = {**scope()["course"], "shortname": "VGA-ES-2023-96-R30"}
    student = academic_student(snapshots=[{"course": item, "error": ""}])
    assert {r["status"] for r in student["rows"]} == {"SIN_PENSUM", "NO_VERIFICABLE"}
    student = academic_student(subjects=[], snapshots=[])
    assert {r["status"] for r in student["rows"]} == {"SIN_PENSUM"}


def academic_service(*, registry=None, enrollments=None, directory=None, course_error=False):
    moodle = MagicMock()
    moodle.get_all_users = AsyncMock(return_value=directory if directory is not None else [user(), user(2), user(3), user(4)])
    moodle.get_user_courses = AsyncMock(side_effect=MoodleApiError("private") if course_error else None, return_value=[scope()["course"]])
    moodle.get_course_enrolled_users = AsyncMock(return_value=[user(), user(2), user(4)])
    base = validation.MoodleEnrollmentValidationService(moodle)
    base._load_academic = MagicMock(return_value={"period": {"code": 10, "name": "PERIODO 10"}, "subjects": [subject()],
        "registry": registry if registry is not None else [person(), person(2), person(3, state="I"), person(4)],
        "enrollments": enrollments if enrollments is not None else [enrollment(), enrollment(2, career=2), enrollment(3), enrollment(4)]})
    base._load_history = MagicMock(return_value=([], [], []))
    base._load_enrollment_headers = MagicMock(return_value=[])
    return academic_validation.MoodleAcademicValidationService(base), moodle


async def finish_academic_scan(service):
    job = await service.start(period_code=10, career_code=1, actor="admin")
    duplicate = await service.start(period_code=10, career_code=1, actor="admin")
    assert duplicate["job_id"] == job["job_id"]
    for task in list(service._tasks):
        await task
    return service.get(job["job_id"], "admin")


def test_academic_scan_is_sequential_active_career_scoped_reuses_aulas_and_owns_reports():
    service, moodle = academic_service()
    job = asyncio.run(finish_academic_scan(service))
    assert job["status"] == "completed" and job["processed"] == job["total"] == 2
    assert [c.args[0] for c in moodle.get_user_courses.await_args_list] == [1, 4]
    assert moodle.get_course_enrolled_users.await_count == 1
    report = job["report"]
    assert [s["student_code"] for s in report["students"]] == ["1", "4"]
    assert report["summary"]["inactive"] == 1
    assert report["summary"]["matches"] == 2
    assert service.validation.report(report["report_id"], "admin") == report
    with pytest.raises(HTTPException):
        service.get(job["job_id"], "other")
    with pytest.raises(HTTPException):
        service.validation.report(report["report_id"], "other")
    sheet = load_workbook(BytesIO(excel_report(report)))["Matriculas academicas"]
    assert sheet.max_row == 3 and sheet["D2"].value == "CARRERA 1" and sheet["F2"].value == "10"
    text = "\n".join(p.extract_text() for p in PdfReader(BytesIO(pdf_report(report))).pages)
    assert "CARRERA 1" in text and "PERIODO 10" in text and "período del aula Moodle" in text


@pytest.mark.parametrize("kwargs, expected", [
    ({"course_error": True}, "NO_VERIFICABLE"),
    ({"directory": []}, "CUENTA_NO_ENCONTRADA"),
    ({"directory": [user(), user(id=44)]}, "IDENTIDAD_AMBIGUA"),
    ({"directory": [user(suspended=True)]}, "SUSPENDIDO_MOODLE"),
])
def test_academic_scan_identity_and_api_errors_never_silently_claim_students_are_correct(kwargs, expected):
    service, _ = academic_service(registry=[person()], enrollments=[enrollment()], **kwargs)
    job = asyncio.run(finish_academic_scan(service))
    assert job["status"] == "completed"
    assert {r["status"] for r in job["report"]["rows"]} == {expected}
    assert "private" not in str(job)


def test_academic_scan_handles_global_errors_and_empty_scope_without_queries():
    service, moodle = academic_service(enrollments=[])
    assert asyncio.run(finish_academic_scan(service))["status"] == "error"
    moodle.get_all_users.assert_not_awaited()
    service, moodle = academic_service()
    moodle.get_all_users.side_effect = MoodleApiError("private upstream")
    job = asyncio.run(finish_academic_scan(service))
    assert job["status"] == "error" and job["report"] is None and "private" not in job["error"]


def test_academic_scan_keeps_active_students_with_header_but_no_subjects():
    service, moodle = academic_service(registry=[person()], enrollments=[])
    service.validation._load_enrollment_headers.return_value = [{"student_code": 1, "career_code": 1, "career": "CARRERA 1", "period_code": 10, "period": "PERIODO 10"}]
    moodle.get_user_courses.return_value = []
    report = asyncio.run(finish_academic_scan(service))["report"]
    assert len(report["students"]) == 1
    assert report["students"][0]["academic_subjects"] == []
    assert report["rows"][0]["status"] == "SIN_MATERIAS_ACADEMICAS"


def test_academic_endpoints_validate_parameters_and_own_permission():
    app = FastAPI()
    app.include_router(router.router)
    service = MagicMock()
    service.start = AsyncMock(return_value={"job_id": "job", "status": "running"})
    service.get.return_value = {"job_id": "job", "status": "completed"}
    app.dependency_overrides[router.get_moodle_academic_validation_service] = lambda: service
    app.dependency_overrides[router._MOODLE_ENROLLMENT_VALIDATION_ACCESS] = lambda: SessionUser(login="admin", rol="ADMINISTRADOR")
    with TestClient(app) as client:
        assert client.post("/api/moodle/enrollment-validation/academic", json={"period_code": 10, "career_code": 1}).status_code == 202
        assert service.start.await_args.kwargs == {"period_code": 10, "career_code": 1, "actor": "admin"}
        assert client.get("/api/moodle/enrollment-validation/academic/job").status_code == 200
        service.get.assert_called_once_with("job", "admin")
        assert client.post("/api/moodle/enrollment-validation/academic", json={"period_code": 10, "career_code": 0}).status_code == 422
        def denied():
            raise HTTPException(403, "No autorizado")
        app.dependency_overrides[router._MOODLE_ENROLLMENT_VALIDATION_ACCESS] = denied
        assert client.get("/api/moodle/enrollment-validation/academic/job").status_code == 403


def selected_report(**values):
    academic = {"period": {"code": 10, "name": "PERIODO 10"}, "registry": [person(), person(2)],
                "subjects": [subject(), subject(1, 102, "VGA-ES-2023-96")],
                "enrollments": [enrollment(), enrollment(number=102), enrollment(2)]}
    academic.update(values.pop("academic", {}))
    return academic_validation.selected_course_audit(academic=academic, headers=values.pop("headers", []),
        career={"code": 1, "name": "CARRERA 1"}, codes=values.pop("codes", [1, 2]),
        people=validation._identity_indexes(academic["registry"])[0], history=values.pop("history", []),
        changes=values.pop("changes", []), snapshots=values.pop("snapshots", [scope(parallel="*")]),
        actor="admin", warnings=[])


def test_unified_selection_compares_both_directions_without_flagging_unselected_subjects():
    result = selected_report()
    assert set(statuses(result)) == {"COINCIDE", "FALTA_MOODLE"}
    assert {row["course_id"] for row in result["rows"]} == {"20"}
    assert [s["student_code"] for s in result["students"]] == ["1", "2"]
    assert result["summary"]["matches"] == result["summary"]["findings"] == 1
    first = result["students"][0]
    assert first["evaluated_subjects"] == 1 and first["findings"] == 0
    assert {s["subject_id"]: (s["in_scope"], s["linked"]) for s in first["academic_subjects"]} == {
        101: (True, True), 102: (False, False)}
    assert result["students"][1]["moodle_courses"] == []
    assert result["course_selection"] is True and result["career"]["code"] == 1
    sheet = load_workbook(BytesIO(excel_report(result)))["Matriculas academicas"]
    assert any(row[-1] == "Fuera de la selección" for row in sheet.iter_rows(min_row=2, values_only=True))
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(pdf_report(result))).pages)
    assert "PERIODO 10" in text and "CARRERA 1" in text and "Falta en Moodle" in text


def test_unified_selection_retains_exact_parallel_and_career_when_a_subject_is_shared():
    result = selected_report(academic={"subjects": [subject(), subject(2, 201)],
        "enrollments": [enrollment(), enrollment(career=2, number=201), enrollment(2, parallel="PB2")]},
        snapshots=[scope([user(), user(2)], parallel="PB1")])
    assert set(statuses(result)) == {"COINCIDE", "OTRO_PARALELO"}
    assert result["students"][0]["academic_subjects"][0]["linked"] is True
    assert result["students"][1]["academic_subjects"][0]["in_scope"] is False
    assert result["students"][1]["findings"] == 1


def test_unified_selection_ignores_other_career_students_but_keeps_unresolved_course_findings():
    result = selected_report(academic={"registry": [person(), person(2), person(3, state="I")],
        "enrollments": [enrollment(), enrollment(2, career=2), enrollment(3)]}, codes=[1],
        snapshots=[scope([user(), user(2), user(3), user(9)], parallel="*")])
    assert {row["student_code"] for row in result["rows"]} == {"1", ""}
    assert result["summary"]["findings"] == 1
    assert any("sin estudiante confirmado" in warning for warning in result["warnings"])


@pytest.mark.parametrize("snapshot, statuses_expected", [
    (scope([], error="No disponible"), {"ERROR_MOODLE", "NO_VERIFICABLE"}),
    (scope([user(username="otro@intec.edu.ec")]), {"IDENTIDAD_AMBIGUA", "NO_VERIFICABLE"}),
    (scope([user(), user(id=8)]), {"DUPLICADO_MOODLE", "NO_VERIFICABLE", "FALTA_MOODLE"}),
])
def test_unified_selection_errors_do_not_report_false_matches(snapshot, statuses_expected):
    result = selected_report(snapshots=[snapshot])
    assert set(statuses(result)) == statuses_expected
    assert result["summary"]["matches"] == 0


def test_unified_selection_handles_headers_without_subjects_and_wrong_career_subjects():
    headers = [{"student_code": 2, "career_code": 1, "career": "CARRERA 1", "period_code": 10}]
    result = selected_report(academic={"subjects": [subject(), subject(2, 201, "VGA-ES-2023-96")],
        "enrollments": [enrollment()]}, headers=headers,
        snapshots=[scope(course={"id": 21, "fullname": "OTRA MATERIA", "shortname": "VGA-ES-2023-96-R30"})])
    assert set(statuses(result)) == {"OTRA_CARRERA", "SIN_MATERIAS_ACADEMICAS"}
    assert result["students"][0]["evaluated_subjects"] == 0
    assert result["students"][1]["academic_subjects"] == []


def test_selected_scan_reads_only_selected_courses_and_preserves_active_scope_and_report_ownership():
    service, moodle = academic_service()
    second_course = {"id": 21, "fullname": "OTRO CURSO", "shortname": "VGA-ES-2023-96-R30"}
    moodle.get_all_courses = AsyncMock(return_value=[scope()["course"], second_course])
    async def run():
        job = await service.start(period_code=10, career_code=1, courses=[{"id": 20, "parallel": "*"}], actor="admin")
        with pytest.raises(HTTPException) as exc:
            await service.start(period_code=10, career_code=1, courses=[{"id": 21, "parallel": "*"}], actor="admin")
        assert exc.value.status_code == 409
        duplicate = await service.start(period_code=10, career_code=1, courses=[{"id": 20, "parallel": "*"}], actor="admin")
        assert duplicate["job_id"] == job["job_id"]
        await asyncio.gather(*list(service._tasks))
        return service.get(job["job_id"], "admin")
    job = asyncio.run(run())
    assert job["status"] == "completed" and job["unit"] == "courses"
    assert job["processed"] == job["total"] == 1
    assert [call.args[0] for call in moodle.get_course_enrolled_users.await_args_list] == [20]
    moodle.get_user_courses.assert_not_awaited()
    moodle.get_all_users.assert_not_awaited()
    result = job["report"]
    assert result["summary"]["inactive"] == 1 and result["summary"]["matches"] == 2
    assert service.validation.report(result["report_id"], "admin") == result


def test_unified_endpoint_rejects_empty_duplicate_and_invalid_selections():
    app = FastAPI()
    app.include_router(router.router)
    service = MagicMock()
    service.start = AsyncMock(return_value={"job_id": "job", "status": "running"})
    app.dependency_overrides[router.get_moodle_academic_validation_service] = lambda: service
    app.dependency_overrides[router._MOODLE_ENROLLMENT_VALIDATION_ACCESS] = lambda: SessionUser(login="admin", rol="ADMINISTRADOR")
    with TestClient(app) as client:
        endpoint = "/api/moodle/enrollment-validation/academic"
        valid = {"period_code": 10, "career_code": 1, "courses": [{"id": 20, "parallel": "pb1"}]}
        assert client.post(endpoint, json=valid).status_code == 202
        assert service.start.await_args.kwargs["courses"] == [{"id": 20, "parallel": "PB1"}]
        for courses in ([], [{"id": 20, "parallel": "*"}] * 2, [{"id": 0, "parallel": "*"}], [{"id": 20, "parallel": ""}]):
            assert client.post(endpoint, json={**valid, "courses": courses}).status_code == 422


def test_selected_scan_keeps_failed_course_distinct_and_continues_without_false_absences():
    service, moodle = academic_service(registry=[person()], enrollments=[enrollment(), enrollment(number=102)])
    service.validation._load_academic.return_value["subjects"].append(subject(1, 102, "VGA-ES-2023-96"))
    second = {"id": 21, "fullname": "OTRO CURSO", "shortname": "VGA-ES-2023-96-R30"}
    moodle.get_all_courses = AsyncMock(return_value=[scope()["course"], second])
    moodle.get_course_enrolled_users.side_effect = [MoodleApiError("private credentials"), [user()]]
    async def run():
        job = await service.start(period_code=10, career_code=1, actor="admin",
                                  courses=[{"id": 21, "parallel": "*"}, {"id": 20, "parallel": "*"}])
        same = await service.start(period_code=10, career_code=1, actor="admin",
                                   courses=[{"id": 20, "parallel": "*"}, {"id": 21, "parallel": "*"}])
        assert same["job_id"] == job["job_id"]
        with pytest.raises(HTTPException):
            await service.start(period_code=10, career_code=1, actor="admin",
                                 courses=[{"id": 20, "parallel": "PB2"}, {"id": 21, "parallel": "*"}])
        await asyncio.gather(*list(service._tasks))
        return service.get(job["job_id"], "admin")
    job = asyncio.run(run())
    assert job["status"] == "completed" and job["processed"] == 2
    assert set(statuses(job["report"])) == {"ERROR_MOODLE", "NO_VERIFICABLE", "COINCIDE"}
    assert "private credentials" not in str(job)
    assert len(job["report"]["scope"]) == 2


def test_selected_scan_rejects_a_course_removed_after_catalog_loading():
    service, moodle = academic_service()
    moodle.get_all_courses = AsyncMock(return_value=[])
    async def run():
        job = await service.start(period_code=10, career_code=1, actor="admin", courses=[{"id": 20, "parallel": "*"}])
        await asyncio.gather(*list(service._tasks))
        return service.get(job["job_id"], "admin")
    job = asyncio.run(run())
    assert job["status"] == "error" and job["report"] is None
    assert "Actualice el catálogo" in job["error"]
    moodle.get_course_enrolled_users.assert_not_awaited()
