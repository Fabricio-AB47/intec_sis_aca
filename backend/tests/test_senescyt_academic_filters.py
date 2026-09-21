from datetime import date
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pandas as pd
import pytest

from app.core.security import SessionUser
from app.routers import senescyt


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(senescyt.router)
    app.dependency_overrides[senescyt._SENESCYT_ACCESS] = lambda: SessionUser(
        login="admin@intec.edu.ec", rol="ADMINISTRADOR",
    )
    with TestClient(app) as value:
        yield value


def test_student_query_scopes_subject_enrollments_before_ranking_and_requires_active_state():
    cutoff = date(2025, 12, 31)
    with patch.object(senescyt, "_read_sql_dataframe", return_value=pd.DataFrame()) as reader:
        result = senescyt._prepare_student_audit_dataframe([1032, 1033, 1032], cutoff)
    sql, params = reader.call_args.args
    assert params == [1032, 1033, cutoff]
    assert "p.cod_periodo IN (?, ?)" in sql
    assert "TRY_CONVERT(date, cx.Fecha_Matricula) <= ?" in sql
    assert sql.index("cx.Fecha_Matricula) <= ?") < sql.index("ROW_NUMBER()")
    assert "GROUP BY cx.codigo_estud, cx.cod_anio_Basica, cx.codigo_periodo" in sql
    assert "PARTITION BY codigo_estud, cod_anio_Basica" in sql
    assert "e.codigo_estud = cne.codigo_estud" in sql
    assert "e.Cedula_Est = cne.Cedula_Est" in sql
    assert "INNER JOIN dbo.ESTADO estado ON e.Estado = estado.IDESTADO" in sql
    assert "estado.IDESTADO)))) = 'A'" in sql
    assert "cne.fecha_matricula AS fechaMatricula" in sql
    assert "e.fechaMatricula" not in sql
    assert result.empty


def test_student_audit_collapses_duplicate_source_rows_without_losing_other_careers():
    columns = [column for column in senescyt._REPORT_COLUMNS if column not in senescyt._NAME_FIELDS]
    row = {column: None for column in columns}
    row.update({
        "codigoEstud": "123", "numeroIdentificacion": "0106889843",
        "Apellidos_nombre": "PEREZ LOPEZ ANA MARIA", "nombreCarrera": "Administracion",
        "fechaMatricula": date(2025, 10, 1),
    })
    raw = pd.DataFrame([row, dict(row), {**row, "nombreCarrera": "Ciberseguridad"}])
    with patch.object(senescyt, "_read_student_audit_source", return_value=raw) as reader:
        result = senescyt._prepare_student_audit_dataframe([1032], date(2025, 12, 31))
    reader.assert_called_once_with([1032], date(2025, 12, 31))
    assert len(result) == 2
    assert not result.duplicated(["numeroIdentificacion", "nombreCarrera"]).any()


def test_teacher_query_uses_assigned_period_and_active_user_with_inclusive_cutoff():
    cutoff = date(2025, 12, 31)
    with patch.object(senescyt, "_read_sql_dataframe", return_value=pd.DataFrame()) as reader:
        result = senescyt._read_teacher_audit_dataframe([1032, 1033], cutoff)
    sql, params = reader.call_args.args
    assert params == [1032, 1033, cutoff, cutoff]
    assert "cd.codigo_periodo = p.cod_periodo" in sql
    assert "p.cod_periodo IN (?, ?)" in sql
    assert "p.fechain <= ?" in sql
    assert "ingreso.fecha IS NULL OR ingreso.fecha <= ?" in sql
    assert "d.fechaIngresoIES" in sql
    assert "FROM dbo.USUARIOS u" in sql
    assert "IN (N'A', N'ACTIVO', N'ACTIVA')" in sql
    assert "SELECT DISTINCT" in sql
    assert result.empty


@pytest.mark.parametrize("target", ["docentes", "estudiantes"])
def test_no_selection_keeps_all_periods_but_does_not_remove_active_validation(target):
    with patch.object(senescyt, "_read_sql_dataframe", return_value=pd.DataFrame()) as reader:
        senescyt._load_senescyt_audit_dataframe(target)
    sql, params = reader.call_args.args
    assert not params
    assert "p.cod_periodo IN" not in sql
    assert " <= ?" not in sql
    assert ("estado.IDESTADO)))) = 'A'" if target == "estudiantes" else "u.Estado") in sql


def test_catalog_includes_historical_periods_even_if_closed(client):
    periods = pd.DataFrame([
        {"codigo_periodo": 1032, "nombre_periodo": "  C2-2025  ", "fecha_inicio": "2025-10-01", "fecha_fin": "2026-03-01"},
        {"codigo_periodo": 901, "nombre_periodo": "C1-2023", "fecha_inicio": None, "fecha_fin": None},
    ])
    with patch.object(senescyt, "_career_catalog", return_value=[]), patch.object(
        senescyt, "_read_sql_dataframe", return_value=periods,
    ) as reader:
        response = client.get("/api/students/senescyt/catalogo")
    assert response.status_code == 200
    assert response.json()["periods"][0]["nombre_periodo"] == "C2-2025"
    assert response.json()["periods"][1]["fecha_inicio"] is None
    assert "Estado" not in reader.call_args.args[0]


@pytest.mark.parametrize("target", ["estudiantes", "docentes"])
def test_report_and_both_downloads_share_filters_and_one_row_per_career(client, target):
    columns = senescyt._STUDENT_AUDIT_COLUMNS if target == "estudiantes" else senescyt._TEACHER_REPORT_COLUMNS
    dataframe = pd.DataFrame([
        {"codigo": "123", "numeroIdentificacion": "0106889843", "nombreCarrera": "Administracion",
         "nombreCompleto": "PEREZ LOPEZ ANA MARIA", "primerNombre": "ANA"},
        {"codigo": "456", "numeroIdentificacion": "0805575735", "nombreCarrera": "Ciberseguridad",
         "nombreCompleto": "OTRA PERSONA", "primerNombre": "OTRA"},
    ])
    params = [("target", target), ("carrera", "Administracion"), ("periodo", 1032),
              ("periodo", 1033), ("periodo", 1032), ("fecha_limite", "2025-12-31")]
    with patch.object(senescyt, "_load_senescyt_audit_dataframe", return_value=(dataframe, columns)) as reader:
        response = client.get("/api/students/senescyt/datos", params=params)
        assert response.status_code == 200, response.text
        body = response.json()
        assert body["period_filter"] == [1032, 1033]
        assert body["cutoff_date"] == "2025-12-31"
        assert body["active_only"] is True
        assert body["summary"]["total_registros"] == 1
        for mode in ("completo", "faltantes"):
            download = client.get("/api/students/senescyt/datos/export", params=params + [("mode", mode)])
            assert download.status_code == 200, download.text
            with ZipFile(BytesIO(download.content)) as archive:
                assert len(archive.namelist()) == 1
                workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])))
                if mode == "completo":
                    assert workbook.active.max_row == 2
                    assert workbook.active.cell(2, 2).value == "0106889843"
                workbook.close()
        assert reader.call_count == 3
        assert all(call.args == (target, [1032, 1033], date(2025, 12, 31)) for call in reader.call_args_list)


@pytest.mark.parametrize("path", ["datos", "datos/export"])
@pytest.mark.parametrize("params", [
    {"periodo": "abc"}, {"periodo": 0}, {"periodo": -1},
    {"periodo": "1);DROP TABLE PERIODO"}, {"fecha_limite": "2025-02-30"},
    {"fecha_limite": "31/12/2025"},
])
def test_invalid_filters_fail_before_querying(client, path, params):
    with patch.object(senescyt, "_load_senescyt_audit_dataframe") as reader:
        response = client.get(f"/api/students/senescyt/{path}", params=params)
    assert response.status_code == 422
    reader.assert_not_called()


@pytest.mark.parametrize("target", ["estudiantes", "docentes"])
def test_empty_scope_does_not_fall_back_to_unfiltered_population(client, target):
    columns = senescyt._STUDENT_AUDIT_COLUMNS if target == "estudiantes" else senescyt._TEACHER_REPORT_COLUMNS
    with patch.object(senescyt, "_load_senescyt_audit_dataframe", return_value=(pd.DataFrame(columns=columns), columns)) as reader:
        response = client.get("/api/students/senescyt/datos", params={"target": target, "periodo": 999999})
    assert response.status_code == 200
    assert response.json()["summary"]["total_registros"] == 0
    reader.assert_called_once_with(target, [999999], None)
