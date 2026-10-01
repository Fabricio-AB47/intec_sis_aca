import asyncio
from io import BytesIO
from unittest.mock import AsyncMock, patch

import pytest
from docx import Document
from fastapi import HTTPException
from pypdf import PdfReader

from app.routers import portal_academico as portal


TEACHER = {"docente": "DOCENTE PRUEBA", "cedula": "TEST", "correo": "docente@example.test"}
META = {"nombre_materia": "Asignatura de prueba", "cod_materia": "TEST", "detalle_periodo": "Período prueba"}
STUDENTS = [
    {"codigo_estud": 1, "nombre_estudiante": "ESTUDIANTE PENDIENTE", "promedio_final": None},
    {"codigo_estud": 2, "nombre_estudiante": "ESTUDIANTE REPROBADO", "promedio_final": 6},
]


@pytest.mark.parametrize("text", ["", "Pendiente", "JUSTIFICACION INICIO " + "seguimiento académico " * 80 + "FIN JUSTIFICACION"])
def test_pdf_always_includes_justification_before_signature(text):
    pdf = portal._teacher_compliance_model_pdf(TEACHER, META, STUDENTS, {}, {"failure_justification": text})
    reader = PdfReader(BytesIO(pdf))
    content = " ".join(" ".join(page.extract_text() for page in reader.pages).split())
    assert content.index("Justificación de estudiantes") < content.rindex("Firma electrónica")
    assert "ESTUDIANTE PENDIENTE" in content and "ESTUDIANTE REPROBADO" in content
    if not text:
        assert "No se registró una justificación adicional" not in content
    elif "FIN JUSTIFICACION" in text:
        assert content.index("FIN JUSTIFICACION") < content.rindex("Firma electrónica")
    else:
        assert text in content
    assert f"Página {len(reader.pages)} de {len(reader.pages)}" in reader.pages[-1].extract_text()


def test_empty_justification_omits_automatic_text_in_pdf_and_docx():
    students = [{"nombre_estudiante": "ESTUDIANTE APROBADO", "promedio_final": 9}]
    assert portal._teacher_compliance_justification_paragraphs(students, {}) == []
    pdf = portal._teacher_compliance_model_pdf(TEACHER, META, students, {}, {})
    pdf_text = " ".join(" ".join(page.extract_text() for page in PdfReader(BytesIO(pdf)).pages).split())
    docx = portal._teacher_compliance_report_docx(TEACHER, META, students, {}, {})
    docx_text = " ".join(p.text for p in Document(BytesIO(docx)).paragraphs)
    for content in (pdf_text, docx_text):
        for removed in (
            "Las observaciones académicas son informativas",
            "No se registran estudiantes reprobados",
            "Justificación del docente:",
            "No se registró una justificación adicional",
        ):
            assert removed not in content
        assert content.index("Justificación de estudiantes") < content.rindex("Firma electrónica")


def test_many_pending_students_are_paginated_without_losing_signature():
    students = [{"nombre_estudiante": f"ESTUDIANTE CASO {i:03}", "promedio_final": None} for i in range(150)]
    pdf = portal._teacher_compliance_model_pdf(TEACHER, META, students, {}, {"failure_justification": "FIN DE CASOS"})
    reader = PdfReader(BytesIO(pdf))
    content = "\n".join(page.extract_text() for page in reader.pages)
    assert "ESTUDIANTE CASO 149" in content
    assert "FIN DE CASOS" in reader.pages[-1].extract_text()
    assert "Firma electrónica" in reader.pages[-1].extract_text()
    assert f"Página {len(reader.pages)} de {len(reader.pages)}" in reader.pages[-1].extract_text()


def test_docx_includes_justification_before_signature():
    data = portal._teacher_compliance_report_docx(TEACHER, META, STUDENTS, {}, {"failure_justification": "JUSTIFICACION DOCUMENTADA"})
    content = "\n".join(p.text for p in Document(BytesIO(data)).paragraphs)
    assert content.index("JUSTIFICACION DOCUMENTADA") < content.index("Firma electrónica")


def test_moodle_outage_is_advisory_and_omits_unverified_resources():
    with patch.object(portal, "_teacher_compliance_moodle_context", AsyncMock(side_effect=HTTPException(503))):
        resources, validation = asyncio.run(portal._prepare_teacher_compliance_generation(
            None, [1055], "TEST", "A", None, [{"course_id": 10, "module_id": 20}], [1], ""))
    assert resources == []
    assert validation["can_generate"]
    assert validation["blockers"]


@pytest.mark.parametrize("status", [400, 401, 403, 404, 500])
def test_access_and_input_errors_are_not_bypassed(status):
    with patch.object(portal, "_teacher_compliance_moodle_context", AsyncMock(side_effect=HTTPException(status))):
        with pytest.raises(HTTPException) as caught:
            asyncio.run(portal._prepare_teacher_compliance_generation(None, [1055], "TEST", "A", None, [], [1], ""))
    assert caught.value.status_code == status


def test_generation_continues_with_missing_grades_and_no_justification():
    context = {"grade_validation": {"can_generate": True, "blockers": ["Notas pendientes"]}}
    with patch.object(portal, "_teacher_compliance_moodle_context", AsyncMock(return_value=context)):
        _, validation = asyncio.run(portal._prepare_teacher_compliance_generation(None, [1055], "TEST", "A", None, [], [1], ""))
    assert validation["can_generate"]
