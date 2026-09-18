import asyncio
from io import BytesIO
import json
from types import SimpleNamespace
from unittest.mock import AsyncMock, MagicMock, Mock, patch
from uuid import uuid4

from fastapi import FastAPI, HTTPException
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pydantic import ValidationError
import pytest

from app.core.security import SessionUser
from app.routers import credential_generator as credentials, direct_admission as admission
from app.services import direct_admission_credentials as service
from app.services import direct_admission_excel as excel


USER = SessionUser(login="admin@example.test", rol="ADMINISTRADOR")
NAMES = service.CredentialNames(primer_nombre="Ana Maria", primer_apellido="De la Cruz", segundo_apellido="Perez")
IDENTITY = "1724036536"
PERSON = {**NAMES.model_dump(), "cedula": IDENTITY, "operator": USER.login, "tipo_persona": "ESTUDIANTE"}
EMAIL = "anamaria.delacruz@intec.edu.ec"
LICENSE = credentials._EducationLicense("ESTUDIANTE", "sku-1", "STANDARDWOFFPACK_STUDENT", "Office 365 A1", "Enabled", 100, 50, 50)
GRAPH = {"id": "graph-1", "employeeId": IDENTITY, "userPrincipalName": EMAIL, "assignedLicenses": [{"skuId": "sku-1"}], "accountEnabled": True}
MOODLE = {"id": 99, "idnumber": IDENTITY, "username": EMAIL, "email": EMAIL, "suspended": False}
REQUEST = uuid4()
PROFILE = {"record": {"identificacion": IDENTITY, "nombre_estudiante": "DE LA CRUZ PEREZ ANA MARIA",
                       "datos_credenciales": {"nombres": "ANA MARIA", "apellidos": "DE LA CRUZ PEREZ"}},
           "identity": IDENTITY, "name": "DE LA CRUZ PEREZ ANA MARIA", "email": "", "code": 123, "period": 1060}
SYNCED_PROFILE = {**PROFILE, "email": EMAIL, "correo_datos": EMAIL, "correo_credenciales": EMAIL,
                  "personal_email": "ana@example.test"}


@pytest.fixture
def remotes():
    client = Mock()
    client.get_users_by_field = AsyncMock(return_value=[])
    client.create_users = AsyncMock(return_value=[{"id": 99, "username": EMAIL}])
    with patch.object(credentials, "_graph_is_configured", return_value=True), \
         patch.object(credentials, "_moodle_is_configured", return_value=True), \
         patch.object(credentials, "get_settings", return_value=SimpleNamespace(moodle_base_url="https://moodle.example.test", moodle_verify_tls=True, moodle_timeout_seconds=30)), \
         patch.object(credentials, "_graph_domain", return_value="intec.edu.ec"), \
         patch.object(credentials, "_existing_email_for_cedula", return_value="") as existing, \
         patch.object(credentials, "_graph_users_by_employee_id", return_value=[]) as graph_id, \
         patch.object(credentials, "_graph_user", return_value=None) as graph_email, \
         patch.object(credentials, "_reserve_identity", side_effect=lambda _person, email, _operator: email) as reserve, \
         patch.object(service, "_local_owner_check") as owner, \
         patch.object(credentials, "_create_graph_user", return_value={"id": "graph-1", "usageLocation": "EC"}) as create_graph, \
         patch.object(credentials, "_assign_graph_license", return_value=("ASIGNADA_ESTUDIANTE", "")) as license_writer:
        yield SimpleNamespace(client=client, existing=existing, graph_id=graph_id, graph_email=graph_email,
                              reserve=reserve, owner=owner, graph_writer=create_graph, license_writer=license_writer)


def resolve(remote, known=""):
    return asyncio.run(service._remote_identity(PERSON, remote.client, known))


def provision(remote, known="", license_info=LICENSE):
    return asyncio.run(service._provision(PERSON, known, remote.client, license_info))


def test_existing_rules_and_compound_names_are_preserved(remotes):
    row = provision(remotes)
    assert row["correo_institucional"] == EMAIL and row["estado_general"] == "COMPLETO"
    assert row["estado_graph"] == "CREADO_GRAPH" and row["estado_moodle"] == "CREADO_MOODLE"
    assert row["clave_permanente"] == credentials._permanent_password(PERSON)
    assert remotes.graph_writer.call_args.args[0]["cedula"] == IDENTITY
    user = remotes.client.create_users.call_args.args[0][0]
    assert user["username"] == EMAIL and user["idnumber"] == IDENTITY
    assert user["firstname"] == "ANA MARIA" and user["lastname"] == "DE LA CRUZ PEREZ"
    assert remotes.license_writer.call_args.args[1] == LICENSE


def test_existing_verified_accounts_are_reused_without_password_reset(remotes):
    remotes.graph_id.return_value = [GRAPH]
    remotes.graph_email.return_value = GRAPH
    remotes.client.get_users_by_field.return_value = [MOODLE]
    row = provision(remotes)
    assert row["estado_general"] == "COMPLETO" and not row["clave_emitida"] and row["clave_permanente"] == ""
    assert "no fue modificada" in row["observacion"]
    remotes.graph_writer.assert_not_called()
    remotes.client.create_users.assert_not_called()


@pytest.mark.parametrize("existing_system", ["office", "moodle"])
def test_only_the_missing_system_account_is_created_for_the_same_cedula(remotes, existing_system):
    if existing_system == "office":
        remotes.graph_id.return_value = [GRAPH]
        remotes.graph_email.return_value = GRAPH
    else:
        remotes.client.get_users_by_field.return_value = [MOODLE]
    row = provision(remotes)
    assert row["estado_general"] == "COMPLETO" and row["correo_institucional"] == EMAIL
    if existing_system == "office":
        remotes.graph_writer.assert_not_called()
        remotes.client.create_users.assert_called_once()
    else:
        remotes.graph_writer.assert_called_once()
        remotes.client.create_users.assert_not_called()


@pytest.mark.parametrize("mutation", [
    {"employeeId": "other"}, {"employeeId": ""}, {"accountEnabled": False}, {"id": ""},
])
def test_unknown_wrong_or_disabled_office_identity_is_never_adopted(remotes, mutation):
    remotes.graph_email.return_value = {**GRAPH, **mutation}
    with pytest.raises(RuntimeError):
        resolve(remotes, EMAIL)
    remotes.graph_writer.assert_not_called()
    remotes.reserve.assert_not_called()


@pytest.mark.parametrize("graph_matches,moodle_matches", [
    ([GRAPH, {**GRAPH, "id": "graph-2"}], []),
    ([], [MOODLE, {**MOODLE, "id": 100}]),
    ([GRAPH], [{**MOODLE, "email": "different@intec.edu.ec"}]),
    ([{**GRAPH, "userPrincipalName": "external@example.test"}], []),
])
def test_remote_duplicates_or_cross_system_email_disagreement_block_before_creation(remotes, graph_matches, moodle_matches):
    remotes.graph_id.return_value = graph_matches
    remotes.client.get_users_by_field.return_value = moodle_matches
    with pytest.raises(RuntimeError):
        resolve(remotes)
    remotes.graph_writer.assert_not_called()
    remotes.client.create_users.assert_not_called()


def test_username_collision_uses_the_established_email_suffix_sequence(remotes):
    async def lookup(field, values):
        return [{**MOODLE, "idnumber": "another"}] if field == "username" and values == [EMAIL] else []
    remotes.client.get_users_by_field.side_effect = lookup
    email, _, _ = resolve(remotes)
    assert email == "anamaria.delacruzp@intec.edu.ec"


def test_duplicate_username_and_email_identities_are_blocked(remotes):
    async def lookup(field, _values):
        return [MOODLE] if field == "email" else [{**MOODLE, "id": 100}] if field == "username" else []
    remotes.client.get_users_by_field.side_effect = lookup
    with pytest.raises(RuntimeError):
        resolve(remotes, EMAIL)
    remotes.reserve.assert_not_called()


def test_local_reservation_is_kept_on_retry_and_cannot_be_replaced(remotes):
    remotes.existing.return_value = "reserved@intec.edu.ec"
    assert resolve(remotes)[0] == "reserved@intec.edu.ec"
    with pytest.raises(RuntimeError):
        resolve(remotes, EMAIL)


def test_same_identity_is_not_silently_reserved_under_a_different_email(remotes):
    remotes.reserve.side_effect = None
    remotes.reserve.return_value = "different@intec.edu.ec"
    with pytest.raises(RuntimeError, match="otro correo"):
        resolve(remotes)


def test_remote_read_failure_does_not_trigger_account_creation(remotes):
    remotes.graph_id.side_effect = TimeoutError("Office timeout")
    with pytest.raises(TimeoutError):
        resolve(remotes)
    remotes.reserve.assert_not_called()
    remotes.client.create_users.assert_not_called()


def test_no_license_capacity_blocks_new_accounts_but_not_already_licensed_accounts(remotes):
    exhausted = credentials._EducationLicense("ESTUDIANTE", "sku-1", "STANDARDWOFFPACK_STUDENT", "A1", "Enabled", 100, 100, 0)
    with pytest.raises(RuntimeError, match="licencias"):
        provision(remotes, license_info=exhausted)
    remotes.graph_writer.assert_not_called()
    remotes.graph_email.return_value = GRAPH
    remotes.graph_id.return_value = [GRAPH]
    remotes.client.get_users_by_field.return_value = [MOODLE]
    assert provision(remotes, license_info=exhausted)["estado_general"] == "COMPLETO"


def test_office_failure_does_not_create_a_disconnected_moodle_account(remotes):
    remotes.graph_writer.side_effect = TimeoutError("Office creation timeout; verify existing account")
    row = provision(remotes)
    assert row["estado_graph"] == "ERROR_GRAPH" and row["estado_general"] == "ERROR"
    assert not row["clave_emitida"]
    remotes.client.create_users.assert_not_called()


def test_license_or_moodle_failure_keeps_confirmed_office_result_and_issued_secret_for_archive(remotes):
    remotes.license_writer.side_effect = RuntimeError("License error")
    remotes.client.create_users.side_effect = RuntimeError("Moodle timeout")
    row = provision(remotes)
    assert row["estado_general"] == "PARCIAL" and row["estado_graph"] == "CREADO_GRAPH"
    assert row["estado_moodle"] == "ERROR_MOODLE" and row["clave_emitida"] and row["clave_permanente"]
    assert "License error" in row["error_licencia"] and "Moodle timeout" in row["error_moodle"]


def test_strict_moodle_recheck_detects_a_late_username_collision(remotes):
    row = asyncio.run(credentials._provision_moodle(remotes.client, PERSON, EMAIL, "not-a-real-password", strict_identity=True))
    assert row[1] == "CREADO_MOODLE"
    remotes.client.create_users.reset_mock()
    remotes.client.get_users_by_field.side_effect = lambda field, _values: [{**MOODLE, "idnumber": "other"}] if field == "username" else []
    row = asyncio.run(credentials._provision_moodle(remotes.client, PERSON, EMAIL, "not-a-real-password", strict_identity=True))
    assert row[1] == "CONFLICTO_MOODLE" and not row[3]
    remotes.client.create_users.assert_not_called()


def test_moodle_creation_must_confirm_a_user_id(remotes):
    remotes.client.create_users.return_value = [{}]
    row = provision(remotes)
    assert row["estado_moodle"] == "ERROR_MOODLE" and row["estado_general"] == "PARCIAL"


def test_passport_identification_is_not_padded_or_stripped(remotes):
    person = {**PERSON, "cedula": "AB123456"}
    asyncio.run(service._remote_identity(person, remotes.client, ""))
    remotes.graph_id.assert_called_once_with("AB123456")
    assert remotes.reserve.call_args.args[0]["cedula"] == "AB123456"


def test_names_must_match_student_and_credentials_cannot_accept_password_fields():
    assert NAMES.matches("ANA MARIA", "DE LA CRUZ PEREZ")
    with pytest.raises(ValidationError):
        service.CredentialNames(**NAMES.model_dump(), password="not-accepted")
    with pytest.raises(ValidationError):
        admission.DirectAdmissionPayload(solicitud_id=uuid4(), estudiante={"identificacion": IDENTITY, "nombres": "Ana", "apellidos": "Perez", "correo": "ana@example.test", "sexo": 1, "etnia": 1, "estado_civil": 1},
                                        matricula={"cod_anio_basica": 7, "codigo_periodo": 1060, "nivel": 1, "materia_codes": [11], "paralelo": "A", "cod_jornada": 1}, credenciales=NAMES)


@pytest.fixture
def state_connection():
    connection = Mock()
    connection.cursor.return_value.fetchone.return_value = [1]
    with patch.object(service, "get_connection", return_value=connection), \
         patch.object(service, "_profile", return_value=PROFILE), \
         patch.object(service, "_read_state", return_value=None) as state_reader:
        yield connection, state_reader


def test_session_lock_blocks_concurrent_provisioning_without_academic_writes(state_connection):
    connection, _ = state_connection
    connection.cursor.return_value.fetchone.return_value = [-1]
    with pytest.raises(HTTPException) as error:
        service._begin(REQUEST, NAMES, USER)
    assert error.value.status_code == 409
    connection.close.assert_called_once()
    sql = " ".join(call.args[0] for call in connection.cursor.return_value.execute.call_args_list)
    assert "Session" in sql and "DATOS_ESTUD SET" not in sql


def test_begin_validates_exact_name_partition_and_preserves_first_attempt_names(state_connection):
    connection, reader = state_connection
    wrong = service.CredentialNames(primer_nombre="ANA", segundo_nombre="MARIA", primer_apellido="DE LA CRUZ PEREZ")
    reader.return_value = {"persona": NAMES.model_dump(), "result": {"estado_general": "PARCIAL"}}
    with pytest.raises(HTTPException, match="distribuci"):
        service._begin(REQUEST, wrong, USER)
    connection.close.assert_called_once()


def test_completed_provisioning_still_verifies_identity_without_recreating_accounts(state_connection, remotes):
    connection, reader = state_connection
    result = {"estado_general": "COMPLETO", "correo_institucional": EMAIL}
    reader.return_value = {"persona": NAMES.model_dump(), "result": result, "report_id": 88}
    remotes.graph_id.return_value = [GRAPH]
    remotes.graph_email.return_value = GRAPH
    remotes.client.get_users_by_field.return_value = [MOODLE]
    with patch.object(credentials, "_education_license", return_value=LICENSE) as license_read, \
         patch.object(credentials, "_record_audit"), patch.object(credentials, "_ensure_tables"), \
         patch.object(service, "_profile", return_value=SYNCED_PROFILE), \
         patch.object(service, "MoodleClient", return_value=remotes.client):
        response = asyncio.run(service.provision_admission_credentials(REQUEST, NAMES, USER))
    assert response["estado_general"] == "COMPLETO" and response["reporte_credencial_id"] == 88
    license_read.assert_called_once()
    remotes.graph_id.assert_called_once_with(IDENTITY)
    remotes.graph_writer.assert_not_called()
    remotes.client.create_users.assert_not_called()
    connection.close.assert_called_once()


def test_second_request_for_same_academic_student_cannot_provision_concurrently(state_connection):
    connection, _ = state_connection
    connection.cursor.return_value.fetchone.side_effect = [[1], [-1]]
    with pytest.raises(HTTPException) as error:
        service._begin(REQUEST, NAMES, USER)
    assert error.value.status_code == 409
    sql = " ".join(str(call.args) for call in connection.cursor.return_value.execute.call_args_list)
    assert "PORTAL_INGRESO_CREDENCIALES_ESTUDIANTE:123" in sql
    assert "INSERT INTO dbo.PORTAL_INGRESO_CREDENCIALES" not in sql
    connection.close.assert_called_once()


def test_academic_identity_is_checked_before_external_account_queries(remotes):
    with patch.object(service, "_begin", side_effect=HTTPException(409, "Identidad académica duplicada")), pytest.raises(HTTPException):
        asyncio.run(service.provision_admission_credentials(REQUEST, NAMES, USER))
    remotes.graph_id.assert_not_called()
    remotes.graph_email.assert_not_called()
    remotes.client.get_users_by_field.assert_not_called()
    remotes.graph_writer.assert_not_called()
    remotes.client.create_users.assert_not_called()


def test_local_email_ownership_accepts_same_cedula_but_rejects_any_other_owner():
    connection = MagicMock()
    connection.__enter__.return_value = connection
    cursor = connection.cursor.return_value
    with patch.object(service, "get_connection", return_value=connection), patch.object(credentials, "_ensure_tables"):
        cursor.fetchall.return_value = [("1724-036536",), (IDENTITY,)]
        service._local_owner_check(EMAIL, IDENTITY)
        cursor.fetchall.return_value = [(IDENTITY,), ("another",)]
        with pytest.raises(RuntimeError, match="otra identidad"):
            service._local_owner_check(EMAIL, IDENTITY)


def test_profile_requires_unique_matching_student_in_academic_database():
    cursor = Mock()
    cursor.fetchone.return_value = [1]
    cursor.fetchall.return_value = [SimpleNamespace(resultado_json=json.dumps(PROFILE["record"]),
        codigo_estud=123, codigo_periodo=1060, Cedula_Est=IDENTITY, Apellidos_nombre=PROFILE["name"],
        Estado="A", correo_datos="", correo_credenciales="")]
    with patch.object(service, "find_academic_student", return_value=SimpleNamespace(codigo_estud=124)) as search:
        with pytest.raises(HTTPException) as error:
            service._profile(cursor, REQUEST)
    assert error.value.status_code == 409
    search.assert_called_once_with(cursor, IDENTITY)


def test_finishing_archives_secret_encrypted_but_never_puts_it_in_state_or_legacy_passwords(state_connection, remotes):
    connection, reader = state_connection
    row = provision(remotes)
    connection.cursor.return_value.fetchone.return_value = [88]
    reader.return_value = {"report_id": 88}
    with patch.object(credentials, "_record_audit") as audit, patch.object(credentials, "_ensure_tables"), \
         patch.object(service, "_profile", return_value=SYNCED_PROFILE):
        result = service._finish(connection, REQUEST, PERSON, row, PROFILE)
    assert result["reporte_credencial_id"] == 88 and "clave_permanente" not in result
    assert audit.call_args.args[2][0]["clave_permanente"]
    statements = connection.cursor.return_value.execute.call_args_list
    assert all("[Password]" not in call.args[0] and "clave =" not in call.args[0] for call in statements)
    update = next(call for call in statements if "SET resultado_json" in call.args[0])
    assert row["clave_permanente"] not in update.args[1]
    assert "CABECERA_MATRICULA" not in " ".join(call.args[0] for call in statements)
    assert result["estado_correo_academico"] == "SINCRONIZADO"
    sync = next(call for call in statements if "UPDATE dbo.CorreosEstudIntec" in call.args[0])
    assert "CorreoPersonal" in sync.args[0] and sync.args[-1] == PROFILE["code"]
    assert any("IF NOT EXISTS" in call.args[0] and "INSERT INTO dbo.CorreosEstudIntec" in call.args[0] for call in statements)


def test_incomplete_email_sync_is_not_reported_as_success(state_connection, remotes):
    connection, _ = state_connection
    row = provision(remotes)
    connection.cursor.return_value.fetchone.return_value = [88]
    with patch.object(credentials, "_record_audit"), patch.object(credentials, "_ensure_tables"), \
         pytest.raises(HTTPException) as error:
        service._finish(connection, REQUEST, PERSON, row, PROFILE)
    assert error.value.status_code == 409
    connection.commit.assert_not_called()


def test_only_admin_can_provision_even_when_academic_role_has_admission_screen(remotes):
    with patch.object(service, "_begin") as begin, pytest.raises(HTTPException) as error:
        asyncio.run(service.provision_admission_credentials(REQUEST, NAMES, SessionUser(login="academic@example.test", rol="ACADEMICO")))
    assert error.value.status_code == 403
    begin.assert_not_called()
    remotes.graph_id.assert_not_called()


@pytest.mark.parametrize("base_url,verify", [("http://moodle.example.test", True), ("https://moodle.example.test", False)])
def test_sensitive_provisioning_requires_https_and_certificate_verification(remotes, state_connection, base_url, verify):
    settings = credentials.get_settings()
    settings.moodle_base_url, settings.moodle_verify_tls = base_url, verify
    with pytest.raises(HTTPException) as error:
        asyncio.run(service.provision_admission_credentials(REQUEST, NAMES, USER))
    assert error.value.status_code == 409
    state_connection[0].cursor.assert_not_called()


def test_remote_errors_cannot_echo_generated_password_to_api_or_state(remotes):
    password = credentials._permanent_password(PERSON)
    remotes.client.create_users.side_effect = RuntimeError(f"Remote error echoes {password}")
    row = provision(remotes)
    assert password not in row["error_moodle"] and "[CREDENCIAL]" in row["error_moodle"]


def test_missing_configuration_blocks_before_state_writes(state_connection):
    connection, _ = state_connection
    with patch.object(credentials, "_graph_is_configured", return_value=False), pytest.raises(HTTPException):
        asyncio.run(service.provision_admission_credentials(REQUEST, NAMES, USER))
    connection.cursor.assert_not_called()


def test_http_credential_routes_are_scoped_and_no_store():
    app = FastAPI()
    app.include_router(admission.router)
    app.dependency_overrides[admission._ACCESS] = lambda: USER
    with TestClient(app) as client, patch.object(admission, "get_admission_credential_state", return_value={"estado_general": "PENDIENTE"}), \
         patch.object(admission, "provision_admission_credentials", new_callable=AsyncMock, return_value={"estado_general": "COMPLETO"}) as writer:
        response = client.get(f"/api/students/ingreso-directo/{REQUEST}/credenciales")
        assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
        response = client.post(f"/api/students/ingreso-directo/{REQUEST}/credenciales", json=NAMES.model_dump())
        assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
        assert writer.call_args.args[0] == REQUEST


def test_legacy_21_column_workbooks_still_work_without_credentials():
    catalogs = {key: [{"value": "1", "label": "value"}] for key in ["tipodocumento", "Sexo", "EstadoCivil", "Etnia"]}
    content = excel.build_template({"cod_anio_basica": "7", "nombre_basica": "Ciberseguridad"}, catalogs, [])
    workbook = load_workbook(BytesIO(content))
    workbook["Estudiantes"].delete_cols(1, workbook["Estudiantes"].max_column)
    workbook["Par\u00e1metros"]["B1"] = excel.LEGACY_VERSION
    for index, (_, label, _, required) in enumerate(excel.COLUMNS, start=1):
        workbook["Estudiantes"].cell(1, index, label + (" *" if required else ""))
    workbook["Estudiantes"].append([IDENTITY, "1", "PEREZ", "ANA", "ana@example.test", "", "1", "1", "1"])
    buffer = BytesIO()
    workbook.save(buffer)
    rows = excel.read_students(buffer.getvalue(), 7, catalogs)
    assert not rows[0]["errores"] and rows[0]["credenciales"]["primer_nombre"] == ""


@pytest.mark.parametrize("change", [
    {"Estado": "I"}, {"Estado": "G"}, {"Cedula_Est": "other"},
    {"correo_datos": EMAIL, "correo_credenciales": "different@intec.edu.ec"},
])
def test_profile_rejects_inactive_students_changed_identity_or_different_local_emails(change):
    cursor = Mock()
    cursor.fetchone.return_value = [1]
    data = {"resultado_json": json.dumps(PROFILE["record"]), "codigo_estud": 123, "codigo_periodo": 1060,
            "Cedula_Est": IDENTITY, "Apellidos_nombre": PROFILE["name"], "Estado": "A", "correo_datos": "", "correo_credenciales": ""}
    cursor.fetchall.return_value = [SimpleNamespace(**{**data, **change})]
    with pytest.raises(HTTPException) as error:
        service._profile(cursor, REQUEST)
    assert error.value.status_code == 409
    assert all("INSERT" not in call.args[0] and "UPDATE" not in call.args[0] for call in cursor.execute.call_args_list)


def test_history_state_is_readonly_and_keeps_report_reference():
    connection = MagicMock()
    connection.__enter__.return_value = connection
    with patch.object(service, "get_connection", return_value=connection), patch.object(service, "_profile", return_value=PROFILE), \
         patch.object(service, "_read_state", return_value={"persona": NAMES.model_dump(), "result": {"estado_general": "COMPLETO"}, "report_id": 88}):
        state = service.get_admission_credential_state(REQUEST, USER)
    assert state["reporte_credencial_id"] == 88 and state["estado_general"] == "COMPLETO"
    connection.cursor.return_value.execute.assert_not_called()


def test_failed_state_persistence_releases_lock_and_preserves_academic_data(state_connection, remotes):
    connection, _ = state_connection
    with patch.object(credentials, "_education_license", return_value=LICENSE), \
         patch.object(service, "_provision", new_callable=AsyncMock, return_value={"estado_general": "COMPLETO"}), \
         patch.object(service, "_finish", side_effect=RuntimeError("State database unavailable")), \
         patch.object(service, "MoodleClient", return_value=remotes.client):
        with pytest.raises(RuntimeError, match="unavailable"):
            asyncio.run(service.provision_admission_credentials(REQUEST, NAMES, USER))
    connection.rollback.assert_called_once()
    connection.close.assert_called_once()
    sql = " ".join(call.args[0] for call in connection.cursor.return_value.execute.call_args_list)
    assert "sp_releaseapplock" in sql and "CABECERA_MATRICULA" not in sql and "CARRERAXESTUD" not in sql
