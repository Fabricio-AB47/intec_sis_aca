from datetime import date
from io import BytesIO
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pandas as pd

from app.core.security import SessionUser
from app.routers import senescyt


HEADERS = """tipoDocumentoId numeroIdentificacion primerApellido segundoApellido primerNombre segundoNombre sexoId generoId estadocivilId etniaId pueblonacionalidadId direccionDomiciliaria provinciaSufragio numeroCelular correoElectronico numDomicilio discapacidad porcentajeDiscapacidad numCarnetDiscapacidad tipoDiscapacidad tipoEnfermedadCatastrofica fechaNacimiento paisNacionalidadId nivelFormacion fechaIngresoIES fechaSalidaIES relacionLaboralIESId ingresoConConcursoMeritos escalafonDocenteId cargoDirectivoId tiempoDedicacionId nombreUnidadAcademica nroasignaturasdocente nroHorasLaborablesSemanaEnCarreraPrograma nroHorasClaseSemanaCarreraPrograma nroHorasInvestigacionSemanaCarreraPrograma nroHorasAdministrativasSemanaCarreraPrograma nroHorasOtrasActividadesSemanaCarreraPrograma nroHorasVinculacionSociedad salarioMensual docenciaTecnicoSuperior docenciaTecnologico estaEnPeriodoSabatico fechaInicioPeriodoSabatico estaCursandoEstudiosId institucionDondeCursaEstudios paisEstudiosId tituloAObtener poseeBecaId tipoBecaId montoBeca financiamientoBecaId pubRevistasCienInIndexadasId numPubRevistasCientifIndexadas docenciaTecnologicoUniversitario docenciaEspecializacionTecnologica docenciaMaestriaTecnologica""".split()


def report(*rows):
    return {"target": "docentes", "dataframe": pd.DataFrame(rows), "report_columns": senescyt._TEACHER_REPORT_COLUMNS}


def test_teacher_model_has_exact_supplied_headers_and_one_row_per_teacher():
    assert len(HEADERS) == 57
    assert senescyt._TEACHER_REPORT_COLUMNS == HEADERS
    teacher = {"codigo": "7", "numeroIdentificacion": "0401105002", "tipoDocumentoId": "1",
               "primerApellido": "GUERRON", "segundoApellido": "MEDINA", "primerNombre": "MARIA",
               "segundoNombre": "FERNANDA", "provinciaSufragio": "05", "numeroCelular": "0984424110",
               "correoElectronico": "maria@intec.edu.ec  ", "fechaNacimiento": date(1974, 5, 20),
               "ingresoConConcursoMeritos": "2", "nombreCarrera": "Administración",
               "nombreCompleto": "GUERRON MEDINA MARIA FERNANDA"}
    content = senescyt._teacher_model_workbook(report(teacher, {**teacher, "nombreCarrera": "Ciberseguridad"}))
    workbook = load_workbook(BytesIO(content), data_only=False)
    assert workbook.sheetnames == ["Sheet1"]
    sheet = workbook.active
    assert sheet.max_row == 2 and sheet.max_column == 57
    assert [cell.value for cell in sheet[1]] == HEADERS
    values = dict(zip(HEADERS, (cell.value for cell in sheet[2])))
    assert values["numeroIdentificacion"] == "0401105002"
    assert values["numeroCelular"] == "0984424110"
    assert values["provinciaSufragio"] == "05"
    assert values["correoElectronico"] == "maria@intec.edu.ec"
    assert values["fechaNacimiento"] == "1974-05-20"
    assert values["ingresoConConcursoMeritos"] == "2"
    assert values["pueblonacionalidadId"] is None
    assert "codigo" not in HEADERS and "nombreCarrera" not in HEADERS
    workbook.close()


def test_teacher_model_is_empty_but_preserves_all_headers_when_no_teachers_exist():
    workbook = load_workbook(BytesIO(senescyt._teacher_model_workbook(report())))
    assert workbook.sheetnames == ["Sheet1"]
    assert workbook.active.max_row == 1 and workbook.active.max_column == 57
    assert [cell.value for cell in workbook.active[1]] == HEADERS
    workbook.close()


def test_teacher_model_does_not_export_a_spreadsheet_formula():
    content = senescyt._teacher_model_workbook(report({"codigo": "1", "primerNombre": "=1+1"}))
    workbook = load_workbook(BytesIO(content), data_only=False)
    cell = workbook.active["E2"]
    assert cell.value == "=1+1" and cell.data_type == "s"
    workbook.close()


def test_teacher_guide_context_and_hours_validation():
    teacher = {"codigo": "7", "numeroIdentificacion": "0401105002", "nombreCarrera": "Administracion",
               "discapacidad": "2", "estaEnPeriodoSabatico": "2", "poseeBecaId": "2",
               "pubRevistasCienInIndexadasId": "2", "fechaIngresoIES": "01/09/2025",
               "provinciaSufragio": 5, "nroHorasLaborablesSemanaEnCarreraPrograma": "40",
               "nroHorasClaseSemanaCarreraPrograma": "24", "nroHorasInvestigacionSemanaCarreraPrograma": "4",
               "nroHorasAdministrativasSemanaCarreraPrograma": "8",
               "nroHorasOtrasActividadesSemanaCarreraPrograma": "2", "nroHorasVinculacionSociedad": "2"}
    workbook = load_workbook(BytesIO(senescyt._teacher_model_workbook(report(teacher))))
    values = dict(zip(HEADERS, (cell.value for cell in workbook.active[2])))
    assert values["fechaIngresoIES"] == "2025-09-01"
    assert values["fechaSalidaIES"] == "NA"
    assert values["fechaInicioPeriodoSabatico"] == "NA"
    assert values["numCarnetDiscapacidad"] == "NA"
    assert values["tipoBecaId"] == 3 and values["financiamientoBecaId"] == 5
    assert values["numPubRevistasCientifIndexadas"] == "NA"
    assert values["provinciaSufragio"] == "05"
    workbook.close()
    assert not senescyt._audit_field_filled(pd.Series({"discapacidad": 1, "numCarnetDiscapacidad": "12345678"}),
                                            "numCarnetDiscapacidad", "docentes")
    assert not senescyt._audit_field_filled(pd.Series({"sexoId": "M"}), "sexoId", "docentes")
    assert senescyt._audit_field_filled(pd.Series(teacher), "nroHorasLaborablesSemanaEnCarreraPrograma", "docentes")
    teacher["nroHorasClaseSemanaCarreraPrograma"] = "20"
    assert not senescyt._audit_field_filled(pd.Series(teacher), "nroHorasLaborablesSemanaEnCarreraPrograma", "docentes")


def test_teacher_query_maps_database_merit_column_to_model_name():
    columns = [column for column in senescyt._TEACHER_REPORT_COLUMNS
               if column not in {"primerApellido", "segundoApellido", "primerNombre", "segundoNombre"}]
    columns += ["codigo", "nombreOriginal", "nombreCarrera"]
    raw = pd.DataFrame([{column: None for column in columns}])
    raw.loc[0, "nombreOriginal"] = "PEREZ LOPEZ ANA MARIA"
    with patch.object(senescyt, "_read_sql_dataframe", return_value=raw) as reader:
        dataframe = senescyt._read_teacher_audit_dataframe()
    sql = reader.call_args.args[0]
    assert "d.ingresoConCursoMeritos AS ingresoConConcursoMeritos" in sql
    assert dataframe.columns.tolist() == HEADERS + ["codigo", "nombreCompleto", "nombreCarrera"]


def test_teacher_complete_download_has_one_model_workbook_per_career():
    app = FastAPI()
    app.include_router(senescyt.router)
    app.dependency_overrides[senescyt._SENESCYT_ACCESS] = lambda: SessionUser(login="admin@intec.edu.ec", rol="ADMINISTRADOR")
    sample = report(
        {"codigo": "1", "numeroIdentificacion": "0401105002", "primerApellido": "PEREZ",
         "primerNombre": "ANA", "nombreCarrera": "Administracion", "fechaIngresoIES": date(2025, 1, 1)},
        {"codigo": "1", "numeroIdentificacion": "0401105002", "primerApellido": "PEREZ",
         "primerNombre": "ANA", "nombreCarrera": "Administracion", "fechaIngresoIES": date(2024, 1, 1)},
        {"codigo": "1", "numeroIdentificacion": "0401105002", "primerApellido": "PEREZ",
         "primerNombre": "ANA", "nombreCarrera": "Ciberseguridad", "fechaIngresoIES": date(2025, 1, 1)},
    )
    with TestClient(app) as client, patch.object(senescyt, "_build_senescyt_audit", return_value=sample):
        response = client.get("/api/students/senescyt/datos/export?target=docentes&mode=completo")
        assert response.status_code == 200
        assert response.headers["content-type"].startswith("application/zip")
        assert ".zip" in response.headers["content-disposition"]
        with ZipFile(BytesIO(response.content)) as archive:
            assert len(archive.namelist()) == 2
            for filename in archive.namelist():
                workbook = load_workbook(BytesIO(archive.read(filename)), read_only=True)
                assert workbook.sheetnames == ["Sheet1"]
                assert workbook.active.max_row == 2
                assert [cell.value for cell in next(workbook.active.iter_rows(max_row=1))] == HEADERS
                workbook.close()


def test_teacher_empty_zip_keeps_model_headers():
    content = senescyt._audit_export_zip(report(), "completo")
    with ZipFile(BytesIO(content)) as archive:
        assert len(archive.namelist()) == 1
        workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])), read_only=True)
        assert [cell.value for cell in next(workbook.active.iter_rows(max_row=1))] == HEADERS
        workbook.close()
