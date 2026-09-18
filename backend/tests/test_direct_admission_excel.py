from datetime import datetime
from io import BytesIO
from unittest.mock import Mock, patch
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import FastAPI, HTTPException, UploadFile
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.core.security import SessionUser
from app.routers import direct_admission as admission
from app.services import direct_admission_excel as excel


USER = SessionUser(login="academic@example.test", rol="ACADEMICO")
CAREER = {"cod_anio_basica": "7", "nombre_basica": "Ciberseguridad"}
CATALOGS = {
    "Sexo": [{"value": "1", "label": "Hombre"}, {"value": "2", "label": "Mujer"}],
    "EstadoCivil": [{"value": "1", "label": "Soltero"}],
    "Etnia": [{"value": "1", "label": "Mestizo"}],
    "tipodocumento": [{"value": "1", "label": "C\u00e9dula"}, {"value": "2", "label": "Pasaporte"}],
    "paisNacionalidadId": [{"value": "1", "label": "Ecuador"}],
    "provinciaNacimeintoId": [{"value": "17", "label": "Pichincha", "parent_value": "1"}],
    "cantonNacimeintoId": [{"value": "1701", "label": "Quito", "parent_value": "17"}],
    "paisResidenciaId": [{"value": "1", "label": "Ecuador"}],
    "codprov": [{"value": "17", "label": "Pichincha", "parent_value": "1"}],
    "Canton": [{"value": "1701", "label": "Quito", "parent_value": "17"}],
}
SUBJECTS = [{"codigo_materia": "21", "nombre_materia": "Segundo", "semestre": 2},
            {"codigo_materia": "11", "nombre_materia": "Primero", "semestre": 1}]
ENROLLMENT = {"cod_anio_basica": 7, "codigo_periodo": 1060, "nivel": 1,
              "materia_codes": [11], "paralelo": "A", "cod_jornada": 1}
STUDENT = {"identificacion": "AB123456", "tipo_documento": "2 - Pasaporte", "apellidos": "Perez Lopez",
           "nombres": "Ana Maria", "correo": "ana@example.test", "sexo": "2 - Mujer", "estado_civil": "1",
           "etnia": "1", "telefono": "023456789", "fecha_nacimiento": "15/06/2000"}


def book(rows=None, change=None, simple=False):
    workbook = load_workbook(BytesIO(excel.build_template(CAREER, CATALOGS, SUBJECTS)))
    sheet = workbook["Estudiantes"]
    columns = excel.SIMPLE_COLUMNS if simple else excel.COLUMNS + excel.CREDENTIAL_COLUMNS
    if not simple:
        sheet.delete_cols(1, sheet.max_column)
        workbook["Par\u00e1metros"]["B1"] = excel.LEGACY_VERSION
        for column, (_, label, _, required) in enumerate(columns, start=1):
            sheet.cell(1, column, label + (" *" if required else ""))
    for number, data in enumerate(rows or [STUDENT], start=2):
        for column, (field, _, _, _) in enumerate(columns, start=1):
            sheet.cell(number, column, data.get(field, ""))
    if change:
        change(workbook)
    result = BytesIO()
    workbook.save(result)
    workbook.close()
    return result.getvalue()


def upload(content, name="estudiantes.xlsx"):
    return UploadFile(filename=name, file=BytesIO(content))


@pytest.fixture
def preview_setup():
    with patch.object(admission, "direct_admission_catalog", return_value={"carreras": [CAREER], "datos_catalogos": CATALOGS}), \
         patch.object(admission, "preview_direct_admission", return_value={"summary": {"insertar": 1}}) as preview:
        yield preview


def validate(content, enrollment=None, create_credentials=False):
    payload = admission.DirectEnrollmentPayload(**(enrollment or ENROLLMENT))
    user = SessionUser(login="admin@example.test", rol="ADMINISTRADOR") if create_credentials else USER
    return admission.validate_direct_admission_excel(user, payload.model_dump_json(), upload(content), create_credentials)


def test_template_contains_career_catalog_dropdowns_and_sorted_pensum():
    workbook = load_workbook(BytesIO(excel.build_template(CAREER, CATALOGS, SUBJECTS)))
    assert workbook.sheetnames == ["Estudiantes", "Par\u00e1metros", "Cat\u00e1logos", "Pensum"]
    assert workbook["Par\u00e1metros"]["B2"].value == "7"
    assert workbook["Par\u00e1metros"]["B3"].value == "Ciberseguridad"
    assert workbook["Pensum"]["A2"].value == "11"
    assert workbook["Estudiantes"].column_dimensions["A"].number_format == "@"
    assert workbook["Estudiantes"].max_column == 6
    assert [cell.value for cell in workbook["Estudiantes"][1]] == [
        "Tipo de documento *", "N\u00famero de c\u00e9dula *", "Primer nombre *",
        "Segundo nombre", "Primer apellido *", "Segundo apellido",
    ]
    validations = list(workbook["Estudiantes"].data_validations.dataValidation)
    assert len(validations) == 1
    assert str(validations[0].sqref) == "A2:A1048576"
    assert all(item.formula1[1:] in workbook.defined_names and item.showErrorMessage for item in validations)
    assert all("1048576" in str(item.sqref) for item in validations)


def test_text_ids_dropdown_codes_date_and_phone_are_preserved():
    student = {**STUDENT, "identificacion": "0805575739", "tipo_documento": "1 - C\u00e9dula"}
    row = excel.read_students(book([student]), 7, CATALOGS)[0]
    assert not row["errores"]
    assert row["datos"]["identificacion"] == "0805575739"
    assert row["datos"]["tipo_documento"] == "1"
    assert row["datos"]["sexo"] == "2"
    assert row["datos"]["telefono"] == "023456789"
    assert row["datos"]["fecha_nacimiento"] == "2000-06-15"


def test_excel_native_date_is_accepted():
    row = excel.read_students(book([{**STUDENT, "fecha_nacimiento": datetime(2000, 6, 15)}]), 7, CATALOGS)[0]
    assert row["datos"]["fecha_nacimiento"] == "2000-06-15" and not row["errores"]


@pytest.mark.parametrize("field,value,fragment", [
    ("identificacion", 805575739, "como texto"),
    ("correo", "=HYPERLINK(\"https://invalid.test\")", "f\u00f3rmulas"),
    ("sexo", "99", "cat\u00e1logo"), ("apellidos", "", "obligatorio"),
    ("fecha_nacimiento", "31/02/2000", "Fecha de nacimiento"),
])
def test_bad_cells_report_row_errors(field, value, fragment):
    row = excel.read_students(book([{**STUDENT, field: value}]), 7, CATALOGS)[0]
    assert row["fila"] == 2 and any(fragment in item for item in row["errores"])


@pytest.mark.parametrize("change,fragment", [
    (lambda workbook: setattr(workbook["Par\u00e1metros"]["B2"], "value", "8"), "carrera"),
    (lambda workbook: setattr(workbook["Par\u00e1metros"]["B1"], "value", "ANTIGUA"), "formato"),
    (lambda workbook: setattr(workbook["Estudiantes"]["B1"], "value", "codigo_estud"), "columnas"),
    (lambda workbook: setattr(workbook["Estudiantes"]["B1"], "value", "Identificaci\u00f3n"), "columnas"),
    (lambda workbook: workbook.remove(workbook["Par\u00e1metros"]), "hojas"),
])
def test_wrong_career_or_changed_structure_is_rejected(change, fragment):
    with pytest.raises(HTTPException) as error:
        excel.read_students(book(change=change), 7, CATALOGS)
    assert error.value.status_code == 400 and fragment in error.value.detail


def test_reordered_known_columns_are_safe():
    def reorder(workbook):
        sheet = workbook["Estudiantes"]
        for row in (1, 2):
            a, b = sheet.cell(row, 1).value, sheet.cell(row, 2).value
            sheet.cell(row, 1, b)
            sheet.cell(row, 2, a)
    data = excel.read_students(book(change=reorder), 7, CATALOGS)[0]
    assert data["datos"]["identificacion"] == "AB123456" and not data["errores"]


def test_extra_nonempty_data_columns_are_not_silently_ignored():
    row = excel.read_students(book(change=lambda workbook: workbook["Estudiantes"].cell(2, 26, "codigo_estud")), 7, CATALOGS)[0]
    assert any("columnas adicionales" in item for item in row["errores"])


@pytest.mark.parametrize("content", [b"", b"not an excel", b"x" * (excel.MAX_FILE_BYTES + 1)], ids=["empty", "invalid", "oversized"])
def test_invalid_or_oversized_files_are_rejected(content):
    with pytest.raises(HTTPException) as error:
        excel.read_students(content, 7, CATALOGS)
    assert error.value.status_code == 400


def test_compressed_bomb_is_rejected_before_openpyxl():
    content = BytesIO()
    with ZipFile(content, "w", ZIP_DEFLATED) as archive:
        archive.writestr("huge.xml", b"x" * (64 * 1024 * 1024 + 1))
    with patch.object(excel, "load_workbook") as reader, pytest.raises(HTTPException):
        excel.read_students(content.getvalue(), 7, CATALOGS)
    reader.assert_not_called()


def test_blank_template_is_rejected():
    with pytest.raises(HTTPException) as error:
        excel.read_students(excel.build_template(CAREER, CATALOGS, SUBJECTS), 7, CATALOGS)
    assert "no contiene" in error.value.detail


def test_valid_rows_get_reviewed_payloads_not_writes(preview_setup):
    with patch.object(admission, "save_direct_admission") as writer:
        result = validate(book())
    assert result["validos"] == 1 and result["invalidos"] == 0
    row = result["items"][0]
    assert row["payload"]["estudiante"]["tipo_documento"] == 2
    assert row["payload"]["estudiante"]["apellidos"] == "PEREZ LOPEZ"
    assert row["payload"]["matricula"]["materia_codes"] == [11]
    assert "codigo_estud" not in row["payload"]["matricula"]
    assert preview_setup.call_count == 1
    writer.assert_not_called()


def test_excel_reuses_a_verified_academic_identity_without_creating_student_during_review(preview_setup):
    preview_setup.return_value = {"summary": {"existentes": 1}, "estudiante": {"accion": "EXISTENTE", "codigo_estud": "123"}}
    with patch.object(admission, "save_direct_admission") as writer:
        result = validate(book())
    row = result["items"][0]
    assert result["validos"] == 1 and row["estudiante_existente"]
    assert row["codigo_estud_existente"] == "123"
    assert "codigo_estud" not in row["payload"]["estudiante"]
    writer.assert_not_called()


@pytest.mark.parametrize("second", [
    {**STUDENT, "correo": "other@example.test"},
    {**STUDENT, "identificacion": "CD123456", "correo": "ANA@EXAMPLE.TEST"},
    {**STUDENT, "identificacion": "CD123456", "correo": "other@example.test", "correo_intec": "ana@example.test"},
])
def test_duplicates_in_any_file_email_or_identity_block_both_rows(preview_setup, second):
    result = validate(book([STUDENT, second]))
    assert result["validos"] == 0 and result["invalidos"] == 2
    assert all(any("filas 2, 3" in item for item in row["errores"]) for row in result["items"])
    preview_setup.assert_not_called()


def test_many_duplicates_have_bounded_messages_without_skipping_validation(preview_setup):
    result = validate(book([STUDENT] * 100))
    assert result["validos"] == 0 and result["invalidos"] == 100
    assert all("80 filas m\u00e1s" in row["errores"][0] and len(row["errores"][0]) < 160 for row in result["items"])
    preview_setup.assert_not_called()


def test_existing_student_is_flagged_and_remaining_rows_still_reviewed(preview_setup):
    preview_setup.side_effect = [HTTPException(409, "Estudiante existente"), {"summary": {"insertar": 1}}]
    result = validate(book([STUDENT, {**STUDENT, "identificacion": "CD123456", "correo": "other@example.test"}]))
    assert result["validos"] == 1 and result["invalidos"] == 1
    assert result["items"][0]["payload"] is None
    assert result["items"][1]["payload"] is not None


def test_prerequisites_and_infrastructure_failures_are_not_hidden(preview_setup):
    preview_setup.return_value = {"summary": {"bloqueadas_por_prerrequisito": 1}}
    result = validate(book())
    assert result["invalidos"] == 1 and "prerrequisitos" in result["items"][0]["errores"][0]
    preview_setup.side_effect = HTTPException(503, "Base temporalmente no disponible")
    with pytest.raises(HTTPException) as error:
        validate(book())
    assert error.value.status_code == 503


def test_model_errors_are_attached_to_exact_row(preview_setup):
    result = validate(book([{**STUDENT, "correo": "invalid email"}]))
    assert result["invalidos"] == 1 and "Correo personal" in result["items"][0]["errores"][0]
    preview_setup.assert_not_called()


def test_more_than_500_rows_are_not_arbitrarily_blocked(preview_setup):
    rows = [{**STUDENT, "identificacion": f"AB{number:06}", "correo": f"student{number}@example.test"} for number in range(501)]
    result = validate(book(rows))
    assert result["total"] == result["validos"] == preview_setup.call_count == 501
    assert len({row["payload"]["solicitud_id"] for row in result["items"]}) == 501


def test_reuses_real_readonly_preview_and_catalog_relationship_validation():
    connection = Mock()
    connection.__enter__ = Mock(return_value=connection)
    connection.__exit__ = Mock(return_value=False)
    with patch.object(admission, "direct_admission_catalog", return_value={"carreras": [CAREER], "datos_catalogos": CATALOGS}), \
         patch.object(admission, "get_connection", return_value=connection), \
         patch.object(admission, "_ensure_new_identity") as identity, \
         patch.object(admission, "_validate_student_catalogs", side_effect=HTTPException(400, "Cant\u00f3n no pertenece a provincia")):
        result = validate(book())
    assert result["invalidos"] == 1 and "provincia" in result["items"][0]["errores"][0]
    connection.rollback.assert_called_once()
    identity.assert_not_called()


def test_http_template_and_multipart_validation(preview_setup):
    app = FastAPI()
    app.include_router(admission.router)
    app.dependency_overrides[admission._ACCESS] = lambda: USER
    with TestClient(app) as client, patch.object(admission, "direct_admission_pensum", return_value={"items": SUBJECTS}):
        response = client.get("/api/students/ingreso-directo/excel/plantilla?cod_anio_basica=7")
        assert response.status_code == 200 and "carrera_7.xlsx" in response.headers["content-disposition"]
        assert load_workbook(BytesIO(response.content))["Par\u00e1metros"]["B2"].value == "7"
        response = client.post("/api/students/ingreso-directo/excel/validar", files={"file": ("students.xlsx", book())},
                               data={"matricula": admission.DirectEnrollmentPayload(**ENROLLMENT).model_dump_json()})
        assert response.status_code == 200 and response.json()["validos"] == 1
        for bad in ["{}", "not json"]:
            response = client.post("/api/students/ingreso-directo/excel/validar", files={"file": ("students.xlsx", book())}, data={"matricula": bad})
            assert response.status_code == 422


def test_inactive_or_wrong_career_is_rejected(preview_setup):
    with pytest.raises(HTTPException) as error:
        validate(book(), {**ENROLLMENT, "cod_anio_basica": 8})
    assert error.value.status_code == 400
    preview_setup.assert_not_called()


def test_excel_credential_names_are_required_only_when_requested_and_do_not_provision_during_review(preview_setup):
    with patch.object(admission, "provision_admission_credentials") as writer:
        invalid = validate(book(), create_credentials=True)
        assert invalid["invalidos"] == 1 and invalid["items"][0]["payload"] is None
        student = {**STUDENT, "primer_nombre": "Ana", "segundo_nombre": "Maria", "primer_apellido": "Perez", "segundo_apellido": "Lopez"}
        valid = validate(book([student]), create_credentials=True)
        assert valid["validos"] == 1
        assert valid["items"][0]["payload"]["credenciales"]["primer_apellido"] == "PEREZ"
        writer.assert_not_called()


def test_office_names_must_match_excel_student_identity(preview_setup):
    student = {**STUDENT, "primer_nombre": "Other", "primer_apellido": "Another"}
    result = validate(book([student]), create_credentials=True)
    assert result["invalidos"] == 1 and "coincidir" in result["items"][0]["errores"][0]
    preview_setup.assert_not_called()


SIMPLE_STUDENT = {"tipo_documento": "1 - C\u00e9dula", "identificacion": "0805575735",
                  "primer_nombre": "Ana", "segundo_nombre": "Mar\u00eda",
                  "primer_apellido": "De la Cruz", "segundo_apellido": "P\u00e9rez"}


def test_six_columns_form_student_identity_and_credential_names_without_personal_email(preview_setup):
    with patch.object(admission, "save_direct_admission") as saver, \
         patch.object(admission, "provision_admission_credentials") as provisioner:
        result = validate(book([SIMPLE_STUDENT], simple=True), create_credentials=True)
    assert result["validos"] == 1 and result["invalidos"] == 0
    row = result["items"][0]
    student = row["payload"]["estudiante"]
    assert student["identificacion"] == "0805575735" and student["tipo_documento"] == 1
    assert student["nombres"] == "ANA MAR\u00cdA"
    assert student["apellidos"] == "DE LA CRUZ P\u00c9REZ"
    assert student["correo"] == student["correo_intec"] == row["correo"] == ""
    assert (student["sexo"], student["estado_civil"], student["etnia"]) == (3, 6, 9)
    assert row["payload"]["credenciales"] == {
        "primer_nombre": "ANA", "segundo_nombre": "MAR\u00cdA",
        "primer_apellido": "DE LA CRUZ", "segundo_apellido": "P\u00c9REZ",
    }
    preview_setup.assert_called_once()
    saver.assert_not_called()
    provisioner.assert_not_called()


def test_optional_second_names_and_surnames_and_unrequested_credentials(preview_setup):
    student = {**SIMPLE_STUDENT, "segundo_nombre": "", "segundo_apellido": ""}
    result = validate(book([student], simple=True))
    payload = result["items"][0]["payload"]
    assert result["validos"] == 1 and payload["credenciales"] is None
    assert payload["estudiante"]["nombres"] == "ANA"
    assert payload["estudiante"]["apellidos"] == "DE LA CRUZ"


@pytest.mark.parametrize("field,value,fragment", [
    ("tipo_documento", "", "obligatorio"),
    ("tipo_documento", "99", "cat\u00e1logo"),
    ("identificacion", 805575739, "como texto"),
    ("identificacion", "0805575738", "c\u00e9dula ecuatoriana"),
    ("primer_nombre", "", "obligatorio"),
    ("primer_apellido", "", "obligatorio"),
    ("segundo_nombre", "=1+1", "f\u00f3rmulas"),
    ("primer_nombre", "123", "nombres y apellidos v\u00e1lidos"),
])
def test_simplified_template_preserves_identity_and_cell_validation(preview_setup, field, value, fragment):
    result = validate(book([{**SIMPLE_STUDENT, field: value}], simple=True), create_credentials=True)
    assert result["invalidos"] == 1 and result["items"][0]["payload"] is None
    assert any(fragment in issue for issue in result["items"][0]["errores"])
    preview_setup.assert_not_called()


def test_missing_personal_emails_do_not_collide_between_different_students(preview_setup):
    other = {**SIMPLE_STUDENT, "tipo_documento": "2 - Pasaporte", "identificacion": "AB123456"}
    result = validate(book([SIMPLE_STUDENT, other], simple=True), create_credentials=True)
    assert result["validos"] == 2 and result["invalidos"] == 0
    assert preview_setup.call_count == 2


def test_repeated_identity_in_simple_file_blocks_every_occurrence(preview_setup):
    result = validate(book([SIMPLE_STUDENT, SIMPLE_STUDENT], simple=True), create_credentials=True)
    assert result["validos"] == 0 and result["invalidos"] == 2
    assert all("filas 2, 3" in row["errores"][0] for row in result["items"])
    preview_setup.assert_not_called()


def test_simplified_excel_reuses_existing_student_by_document(preview_setup):
    preview_setup.return_value = {"summary": {"existentes": 1}, "estudiante": {"accion": "EXISTENTE", "codigo_estud": "123"}}
    result = validate(book([SIMPLE_STUDENT], simple=True), create_credentials=True)
    assert result["validos"] == 1 and result["items"][0]["estudiante_existente"]
    assert result["items"][0]["codigo_estud_existente"] == "123"


def test_simple_header_reordering_is_safe():
    def reorder(workbook):
        sheet = workbook["Estudiantes"]
        for row in (1, 2):
            first, last = sheet.cell(row, 1).value, sheet.cell(row, 6).value
            sheet.cell(row, 1, last)
            sheet.cell(row, 6, first)
    row = excel.read_students(book([SIMPLE_STUDENT], simple=True, change=reorder), 7, CATALOGS)[0]
    assert not row["errores"] and row["datos"]["tipo_documento"] == "1"
    assert row["credenciales"]["segundo_apellido"] == "P\u00e9rez"


def test_simple_format_cannot_add_email_columns():
    def add_column(workbook):
        workbook["Estudiantes"].cell(1, 7, "Correo personal")
    with pytest.raises(HTTPException) as error:
        excel.read_students(book([SIMPLE_STUDENT], simple=True, change=add_column), 7, CATALOGS)
    assert error.value.status_code == 400 and "columnas" in error.value.detail


def test_http_six_column_upload_requires_only_identity_and_screen_enrollment(preview_setup):
    app = FastAPI()
    app.include_router(admission.router)
    app.dependency_overrides[admission._ACCESS] = lambda: SessionUser(login="admin@example.test", rol="ADMINISTRADOR")
    with TestClient(app) as client:
        response = client.post("/api/students/ingreso-directo/excel/validar",
            files={"file": ("students.xlsx", book([SIMPLE_STUDENT], simple=True))},
            data={"matricula": admission.DirectEnrollmentPayload(**ENROLLMENT).model_dump_json(), "crear_credenciales": "true"})
    assert response.status_code == 200 and response.json()["validos"] == 1
    payload = response.json()["items"][0]["payload"]
    assert payload["credenciales"]["primer_nombre"] == "ANA"
    assert payload["matricula"]["cod_anio_basica"] == 7
