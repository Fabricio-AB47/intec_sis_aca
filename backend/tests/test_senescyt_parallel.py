from datetime import date
from io import BytesIO
import os
from unittest.mock import patch
from zipfile import ZipFile

from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pandas as pd
import pytest

from app.core.security import SessionUser
from app.routers import senescyt


def student_source(parallel):
    row = {field: None for field in senescyt._REPORT_COLUMNS if field not in senescyt._NAME_FIELDS}
    row.update({
        "codigoEstud": "123", "tipoDocumentoId": 1, "numeroIdentificacion": "0106889843",
        "Apellidos_nombre": "PEREZ LOPEZ ANA MARIA", "nombreCarrera": "Administracion",
        "fechaMatricula": date(2025, 10, 1), "nivelAcademicoQueCursa": 4,
        "paraleloId": parallel,
    })
    return pd.DataFrame([row])


@pytest.fixture
def client():
    app = FastAPI()
    app.include_router(senescyt.router)
    app.dependency_overrides[senescyt._SENESCYT_ACCESS] = lambda: SessionUser(
        login="admin@intec.edu.ec", rol="ADMINISTRADOR",
    )
    with TestClient(app) as instance:
        yield instance


def test_parallel_comes_from_unique_period_assignment_and_exact_enrollment_catalog_entry():
    cutoff = date(2025, 12, 31)
    with patch.object(senescyt, "_read_sql_dataframe", return_value=pd.DataFrame()) as reader:
        senescyt._read_student_audit_source([1032, 1033, 1032], cutoff)
    sql, params = reader.call_args.args
    assert params == [1032, 1033, cutoff]
    assert "INNER JOIN dbo.PERIODO p ON cx.codigo_periodo = p.cod_periodo" in sql
    assert "p.cod_periodo IN (?, ?)" in sql
    assert sql.index("cx.Fecha_Matricula) <= ?") < sql.index("GROUP BY cx.codigo_estud")
    assert "GROUP BY cx.codigo_estud, cx.cod_anio_Basica, cx.codigo_periodo" in sql
    assert "COUNT(DISTINCT NULLIF(UPPER(LTRIM(RTRIM(cx.paralelo))), '')) = 1" in sql
    assert "COUNT(NULLIF(LTRIM(RTRIM(cx.paralelo)), '')) = COUNT(*)" in sql
    assert "ORDER BY COALESCE(inicio_periodo, fecha_matricula) DESC" in sql
    assert "WHERE s.posicion = 1" in sql
    assert "FROM dbo.PARALELOS par" in sql
    assert "UPPER(LTRIM(RTRIM(par.paralelo))) = s.paralelo" in sql
    assert "TRY_CONVERT(int, par.num) BETWEEN 1 AND 20" in sql
    assert "CASE WHEN COUNT(*) = 1" in sql
    assert "cne.paralelo_id AS paraleloId" in sql
    assert "e.Paralelo" not in sql
    assert "dbo.Paralelo " not in sql
    assert "par.codigo_paralelo" not in sql
    assert "par.activo" not in sql
    assert "LIKE" not in sql


@pytest.mark.parametrize("parallel", [1, 2, 12, 15, 20, None])
def test_preview_and_both_excel_downloads_use_computed_parallel(client, parallel):
    params = [("target", "estudiantes"), ("carrera", "Administracion"),
              ("periodo", 1032), ("periodo", 1033), ("fecha_limite", "2025-12-31")]
    with patch.object(senescyt, "_read_sql_dataframe", return_value=student_source(parallel)) as reader, patch.object(
        senescyt, "_read_dataframe", side_effect=AssertionError("Stored parallel must not be used in reports"),
    ):
        response = client.get("/api/students/senescyt/datos", params=params)
        assert response.status_code == 200, response.text
        row = response.json()["rows"][0]
        assert row["fields"]["paraleloId"] == parallel
        assert ("paraleloId" in row["campos_faltantes"]) == (parallel is None)
        for mode in ("completo", "faltantes"):
            response = client.get("/api/students/senescyt/datos/export", params=params + [("mode", mode)])
            assert response.status_code == 200, response.text
            with ZipFile(BytesIO(response.content)) as archive:
                workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])))
                matched = 0
                for sheet in workbook:
                    headers = [cell.value for cell in sheet[1]]
                    if "paraleloId" in headers:
                        assert sheet.cell(2, headers.index("paraleloId") + 1).value == parallel
                        matched += 1
                assert matched
                workbook.close()
        assert reader.call_count == 3
        assert all(call.args[1] == [1032, 1033, date(2025, 12, 31)] for call in reader.call_args_list)


def test_parallel_is_recalculated_on_each_generation_and_kept_separate_by_career(client):
    first_source = pd.concat([
        student_source(1), student_source(2).assign(nombreCarrera="Ciberseguridad"),
    ], ignore_index=True)
    with patch.object(senescyt, "_read_sql_dataframe", side_effect=[first_source, student_source(3)]):
        first = client.get("/api/students/senescyt/datos?target=estudiantes&periodo=1032")
        second = client.get("/api/students/senescyt/datos?target=estudiantes&periodo=1032")
    assert first.status_code == second.status_code == 200
    parallels = {row["nombre_carrera"]: row["fields"]["paraleloId"] for row in first.json()["rows"]}
    assert parallels == {"Administracion": 1, "Ciberseguridad": 2}
    assert second.json()["rows"][0]["fields"]["paraleloId"] == 3


@pytest.mark.parametrize("parallel, valid", [
    (1, True), (20, True), (None, False), (0, False), (-1, False), (21, False), (1.5, False), ("PB1", False),
])
def test_parallel_must_be_a_valid_guide_code(parallel, valid):
    assert senescyt._audit_field_filled(pd.Series({"paraleloId": parallel}), "paraleloId", "estudiantes") is valid


def test_legacy_download_uses_computed_parallel_without_changing_raw_student_editor(client):
    with patch.object(senescyt, "_read_sql_dataframe", return_value=student_source(20)), patch.object(
        senescyt, "_count_scalar", return_value=1,
    ), patch.object(senescyt, "_read_dataframe", side_effect=AssertionError("Stored parallel must not be exported")):
        response = client.get("/api/students/senescyt/estudiantes/export")
    assert response.status_code == 200, response.text
    with ZipFile(BytesIO(response.content)) as archive:
        workbook = load_workbook(BytesIO(archive.read(archive.namelist()[0])))
        headers = [cell.value for cell in workbook.active[1]]
        assert workbook.active.cell(2, headers.index("paraleloId") + 1).value == 20
        workbook.close()
    assert "e.Paralelo AS paraleloId" in senescyt._QUERY


_SQL_FIXTURES = """
source_periodos AS (
    SELECT * FROM (VALUES
        (101, 'R', CAST('2024-01-01' AS date)),
        (102, 'R', CAST('2025-01-01' AS date)),
        (103, 'H', CAST('2026-01-01' AS date)),
        (104, 'X', CAST('2025-01-01' AS date))
    ) v(cod_periodo, TipoMatricula, fechain)
), source_matriculas AS (
    SELECT * FROM (VALUES
        (1, 1, 101, 1, 'PB1', '2024-01-02'),
        (1, 1, 102, 2, 'A', '2025-01-02'),
        (1, 1, 102, 3, 'PB1', '2025-11-01'),
        (1, 1, 103, 3, 'PB2', '2026-01-02'),
        (1, 2, 102, 4, 'PB1', '2025-01-02'),
        (2, 1, 102, 1, 'A', '2025-01-02'),
        (2, 1, 102, 2, 'PB1', '2025-01-02'),
        (3, 1, 102, 1, ' a ', '2025-01-02'),
        (3, 1, 102, 2, 'A', '2025-01-02'),
        (4, 1, 102, 1, 'PB1', '2025-01-02'),
        (5, 1, 102, 1, NULL, '2025-01-02'),
        (6, 1, 102, 1, '', '2025-01-02'),
        (6, 1, 102, 2, 'A', '2025-01-02'),
        (7, 1, 102, 1, 'N0', '2025-01-02'),
        (8, 1, 102, 1, 'E', '2025-01-02'),
        (9, 1, 102, 1, 'F', '2025-01-02'),
        (10, 1, 102, 1, 'T', '2025-01-02'),
        (11, 1, 102, 1, 'Z', '2025-01-02'),
        (12, 1, 102, 1, 'A', '2025-01-02'),
        (13, 12, 102, 1, 'A', '2025-01-02'),
        (14, 1, 104, 1, 'A', '2025-01-02'),
        (15, 1, 102, 1, 'A', '2025-11-01'),
        (16, 1, 102, 1, 'PBS1', '2025-01-02'),
        (17, 1, 101, 1, 'A', '2024-01-02'),
        (17, 1, 102, 1, 'Z', '2025-01-02'),
        (18, 1, 102, 4, 'A', '2025-01-02'),
        (19, 1, 102, 1, ' pb1', '2025-01-02'),
        (20, 1, 102, 1, 'PBS4', '2025-01-02'),
        (21, 1, 102, 1, 'PB9', '2025-01-02')
    ) v(codigo_estud, cod_anio_Basica, codigo_periodo, codigo_materia, paralelo, Fecha_Matricula)
), source_pensum AS (
    SELECT * FROM (VALUES (1, 1, 1), (1, 2, 4), (1, 3, 3), (2, 4, 9))
    v(Cod_AnioBasica, codigo_materia, Semestre)
), source_estudiantes AS (
    SELECT DISTINCT codigo_estud, CONVERT(varchar(20), codigo_estud) AS Cedula_Est,
        'TEST' AS Apellidos_nombre, CASE WHEN codigo_estud = 12 THEN 'I' ELSE 'A' END AS Estado
    FROM source_matriculas
), source_estados AS (
    SELECT * FROM (VALUES ('A'), ('I')) v(IDESTADO)
), source_carreras AS (
    SELECT * FROM (VALUES (1, 'Carrera 1'), (2, 'Carrera 2'), (12, 'Excluida'))
    v(Cod_AnioBasica, Nombre_Basica)
), source_paralelos AS (
    SELECT * FROM (VALUES
        (1, 'A'), (2, 'PB1'), (3, 'PB2'), (0, 'N0'),
        (21, 'E'), (6, 'F'), (7, ' F '), (20, 'T'), (12, 'PBS1'), (15, 'PBS4')
    ) v(num, paralelo)
)
"""


@pytest.mark.skipif(os.environ.get("SENESCYT_SQL_READONLY_TESTS") != "1", reason="Requires SQL Server; SELECT-only fixtures")
def test_sql_server_period_resolution_with_select_only_fixtures():
    def resolve(periods, cutoff=None):
        with patch.object(senescyt, "_read_sql_dataframe", return_value=pd.DataFrame()) as reader:
            senescyt._read_student_audit_source(periods, cutoff)
        sql, params = reader.call_args.args
        # Exercise the actual report CTEs with synthetic rows, without reading or writing student tables.
        head, separator, _ = sql.partition("\nSELECT\n")
        assert separator and head.lstrip().startswith("WITH matriculas AS (")
        sql = "WITH " + _SQL_FIXTURES + ", " + head.strip()[len("WITH "):]
        tables = {
            "CARRERAXESTUD": "source_matriculas", "PERIODO": "source_periodos",
            "PENSUM": "source_pensum", "DATOS_ESTUD": "source_estudiantes",
            "ESTADO": "source_estados", "CARRERAS": "source_carreras", "PARALELOS": "source_paralelos",
        }
        for table, fixture in tables.items():
            sql = sql.replace(f"dbo.{table}", fixture)
        assert "dbo." not in sql
        sql += " SELECT codigo_estud, nombre_carrera, paralelo_id, nivel_academico FROM matricula_cne_catalogada"
        frame = senescyt._read_sql_dataframe(sql, params)
        return {
            (int(row.codigo_estud), row.nombre_carrera): (
                None if pd.isna(row.paralelo_id) else int(row.paralelo_id),
                None if pd.isna(row.nivel_academico) else int(row.nivel_academico),
            ) for row in frame.itertuples()
        }

    scoped = resolve([102], date(2025, 6, 30))
    assert scoped == {
        (1, "Carrera 1"): (1, 4), (1, "Carrera 2"): (2, 9),
        (2, "Carrera 1"): (None, 4), (3, "Carrera 1"): (1, 4),
        (4, "Carrera 1"): (2, 1), (5, "Carrera 1"): (None, 1),
        (6, "Carrera 1"): (None, 4), (7, "Carrera 1"): (None, 1),
        (8, "Carrera 1"): (None, 1), (9, "Carrera 1"): (None, 1),
        (10, "Carrera 1"): (20, 1), (11, "Carrera 1"): (None, 1),
        (16, "Carrera 1"): (12, 1), (17, "Carrera 1"): (None, 1),
        (18, "Carrera 1"): (1, None),
        (19, "Carrera 1"): (2, 1), (20, "Carrera 1"): (15, 1), (21, "Carrera 1"): (None, 1),
    }
    assert resolve([101]) == {(1, "Carrera 1"): (2, 1), (17, "Carrera 1"): (1, 1)}
    assert resolve([101, 102], date(2025, 6, 30)) == scoped
    assert resolve([101, 102])[(1, "Carrera 1")] == (None, 4)
    assert resolve(None)[(1, "Carrera 1")] == (3, 4)
    assert resolve([999999]) == {}
