from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import LongTable, Paragraph, SimpleDocTemplate, Spacer, TableStyle


REPORT_COLUMNS = [
    ("status_label", "Resultado"), ("student_code", "Código estudiante"), ("student", "Estudiante"),
    ("document", "Identificación"), ("email", "Correo académico"), ("moodle_email", "Correo Moodle"),
    ("moodle_user_id", "ID usuario Moodle"), ("course_id", "ID curso Moodle"), ("course", "Curso Moodle"),
    ("moodle_code", "Código Moodle"), ("subject_code", "Código único materia"), ("subject", "Materia"),
    ("period_code", "Código período"), ("period", "Período"), ("academic_career", "Carrera de matrícula"),
    ("academic_parallel", "Paralelo académico"), ("moodle_parallel", "Paralelo consultado"),
    ("period_careers", "Carreras del período"), ("previous_careers", "Otras carreras históricas"),
    ("career_changes", "Cambios de carrera aplicados"), ("reason", "Observación"),
    ("moodle_idnumber", "Número ID curso Moodle"), ("code_match", "Comparación del código"),
    ("code_match_detail", "Criterio de comparación"), ("code_candidates", "Códigos PENSUM considerados"),
    ("subject_careers", "Carreras de la materia"), ("enrollment_periods", "Matrículas académicas relacionadas"),
    ("period_relation", "Relación de período con Moodle"),
]
SCOPE_TEXT = "Comparación de solo lectura limitada al período y aulas seleccionados. * incluye todos los paralelos. Los hallazgos requieren revisión; no se modificaron matrículas."


def excel_report(report: dict) -> bytes:
    book = Workbook()
    summary = book.active
    summary.title = "Resumen"
    for row in [
        ["Validación de matrículas académicas y Moodle"],
        ["Período", str(report["period"]["code"]), report["period"]["name"]],
        ["Generado", report["generated_at"]], ["Responsable", report["actor"]],
        ["Alcance", report.get("scope_text", SCOPE_TEXT)], ["Registros", report["summary"]["rows"]],
        *([["Carrera", report["career"]["name"], str(report["career"]["code"])]] if report.get("career") else []),
        ["Coincidencias", report["summary"]["matches"]], ["Hallazgos", report["summary"]["findings"]],
        ["Inactivos excluidos", report["summary"]["inactive"]],
        *[["Advertencia", warning] for warning in report["warnings"]],
        ["ID curso", "Curso", "Paralelo", "Código único", "Error de consulta", "Comparación del código", "Criterio"],
        *[[str(c["id"]), c["name"], c["parallel"], c["subject_code"], c["error"],
           c.get("code_match", ""), c.get("code_match_detail", "")] for c in report["scope"]],
    ]:
        summary.append(row)
    detail = book.create_sheet("Detalle")
    detail.append([label for _, label in REPORT_COLUMNS])
    for row in report["rows"]:
        detail.append([str(row.get(key, "")) for key, _ in REPORT_COLUMNS])
    detail.freeze_panes = "A2"
    detail.auto_filter.ref = detail.dimensions
    if "students" in report:
        academic = book.create_sheet("Matriculas academicas")
        academic.append(["Código estudiante", "Estudiante", "Identificación", "Carrera", "Período", "Código período",
                         "Materia", "Código único", "Semestre", "Paralelo", "Materia en aula activa"])
        for student in report["students"]:
            for subject in student["academic_subjects"]:
                academic.append([student["student_code"], student["student"], student["document"], student["academic_career"],
                                 student["period"], student["period_code"], subject["subject"], subject["subject_code"],
                                 subject["semester"], subject["parallel"], "Fuera de la selección" if subject.get("in_scope") is False
                                 else "Sí" if subject["linked"] else "Pendiente de verificación"])
        academic.freeze_panes = "A2"
        academic.auto_filter.ref = academic.dimensions
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Names and identifiers must never become spreadsheet formulas.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="176E7C")
        for index in range(1, sheet.max_column + 1):
            sheet.column_dimensions[get_column_letter(index)].width = 65 if sheet.cell(1, index).value in {"Observación", "Criterio de comparación"} else 32
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def pdf_report(report: dict) -> bytes:
    output = BytesIO()
    style = ParagraphStyle("body", fontName="Helvetica", fontSize=7, leading=10, wordWrap="CJK")
    title = ParagraphStyle("title", parent=style, fontSize=14, leading=18, spaceAfter=8)

    def paragraph(value) -> Paragraph:
        return Paragraph(escape(str(value or "")).replace("\n", "<br/>"), style)

    content = [Paragraph("Validación de matrículas académicas y Moodle", title),
               paragraph(f"Período: {report['period']['code']} - {report['period']['name']}"),
               paragraph(f"Generado: {report['generated_at']} | Responsable: {report['actor']}"),
               paragraph(report.get("scope_text", SCOPE_TEXT)),
               *([paragraph(f"Carrera: {report['career']['name']} ({report['career']['code']})")] if report.get("career") else []),
               paragraph(f"Coincidencias: {report['summary']['matches']} | Hallazgos: {report['summary']['findings']} | Inactivos: {report['summary']['inactive']}"),
               *[paragraph("Advertencia: " + w) for w in report["warnings"]], Spacer(1, 8)]
    for course in report["scope"]:
        content.append(paragraph(f"Aula {course['id']}: {course['name']} | Paralelo: {course['parallel']} | Código: {course['subject_code']} | {course['error']}"))
        content.append(paragraph(f"Comparación: {course.get('code_match', '')}. {course.get('code_match_detail', '')}"))
    content.append(Spacer(1, 10))
    data = [[paragraph(s) for s in ["Resultado", "Estudiante / identidad", "Aula / materia", "Carrera / paralelo", "Antecedentes", "Observación"]]]
    for row in report["rows"]:
        data.append([paragraph(row["status_label"]),
            paragraph(f"{row['student']}\nCódigo: {row['student_code']}\n{row['document']}\n{row['email']}\nMoodle {row['moodle_user_id']}: {row['moodle_email']}"),
            paragraph(f"{row['course']}\nID: {row['course_id']}\nNombre corto: {row['moodle_code']}\nNúmero ID: {row.get('moodle_idnumber', '')}\nPENSUM: {row['subject_code']}\n{row['subject']}"),
            paragraph(f"{row['academic_career']}\nCarreras del período: {row['period_careers']}\nParalelo académico: {row['academic_parallel']}\nConsultado: {row['moodle_parallel']}\nCarreras de la materia: {row.get('subject_careers', '')}\nMatrículas relacionadas: {row.get('enrollment_periods', '')}"),
            paragraph(f"{row['previous_careers']}\n{row['career_changes']}"),
            paragraph(f"{row['reason']}\nComparación: {row.get('code_match', '')}\n{row.get('code_match_detail', '')}\nCandidatos: {row.get('code_candidates', '')}\n{row.get('period_relation', '')}")])
    if not report["rows"]:
        content.append(paragraph("No se encontraron registros comparables en el alcance seleccionado."))
    table = LongTable(data, colWidths=[76, 137, 147, 123, 123, 180], repeatRows=1, splitInRow=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7f2f4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#bbcbd0")),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
    ]))
    content.append(table)

    def page_number(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(810, 15, f"Página {document.page}")

    SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=28, rightMargin=28,
                      topMargin=26, bottomMargin=30).build(content, onFirstPage=page_number, onLaterPages=page_number)
    return output.getvalue()
