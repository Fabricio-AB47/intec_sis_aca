"""Career-bound workbooks for direct admission; no database writes."""

from datetime import date, datetime
from io import BytesIO
import unicodedata
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile, ZipFile

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation


VERSION = "INTEC_INGRESO_DIRECTO_2"
LEGACY_VERSION = "INTEC_INGRESO_DIRECTO_1"
MAX_FILE_BYTES = 12 * 1024 * 1024
# Columns are an explicit whitelist, not arbitrary DATOS_ESTUD attributes.
COLUMNS = [
    ("identificacion", "Identificación", None, True),
    ("tipo_documento", "Tipo de documento", "tipodocumento", True),
    ("apellidos", "Apellidos", None, True),
    ("nombres", "Nombres", None, True),
    ("correo", "Correo personal", None, True),
    ("correo_intec", "Correo institucional", None, False),
    ("sexo", "Sexo", "Sexo", True),
    ("estado_civil", "Estado civil", "EstadoCivil", True),
    ("etnia", "Etnia", "Etnia", True),
    ("fecha_nacimiento", "Fecha de nacimiento", None, False),
    ("telefono", "Teléfono", None, False),
    ("movil", "Celular", None, False),
    ("pais_nacionalidad", "País de nacionalidad", "paisNacionalidadId", False),
    ("provincia_nacimiento", "Provincia de nacimiento", "provinciaNacimeintoId", False),
    ("canton_nacimiento", "Cantón de nacimiento", "cantonNacimeintoId", False),
    ("pais_residencia", "País de residencia", "paisResidenciaId", False),
    ("provincia_residencia", "Provincia de residencia", "codprov", False),
    ("canton_residencia", "Cantón de residencia", "Canton", False),
    ("direccion", "Dirección", None, False),
    ("colegio", "Colegio", None, False),
    ("titulo_bachiller", "Título de bachiller", None, False),
]
CREDENTIAL_COLUMNS = [
    ("primer_nombre", "Primer nombre Office 365", None, False),
    ("segundo_nombre", "Segundo nombre Office 365", None, False),
    ("primer_apellido", "Primer apellido Office 365", None, False),
    ("segundo_apellido", "Segundo apellido Office 365", None, False),
]
SIMPLE_COLUMNS = [
    ("tipo_documento", "Tipo de documento", "tipodocumento", True),
    ("identificacion", "Número de cédula", None, True),
    ("primer_nombre", "Primer nombre", None, True),
    ("segundo_nombre", "Segundo nombre", None, False),
    ("primer_apellido", "Primer apellido", None, True),
    ("segundo_apellido", "Segundo apellido", None, False),
]


def _text_cell(sheet, row, column, value):
    cell = sheet.cell(row, column, str(value))
    cell.data_type = "s"
    cell.number_format = "@"
    return cell


def _options(catalogs, key):
    options = catalogs.get(key, [])
    return [option for option in options if key != "tipodocumento" or str(option["value"]) in {"1", "2", "3"}]


def build_template(career: dict, catalogs: dict, subjects: list[dict]) -> bytes:
    workbook = Workbook()
    students = workbook.active
    students.title = "Estudiantes"
    students.freeze_panes = "C2"
    parameters = workbook.create_sheet("Parámetros")
    for row, (key, value) in enumerate([
        ("Formato", VERSION), ("Código de carrera", career["cod_anio_basica"]),
        ("Carrera", career["nombre_basica"]),
        ("Ingreso", "Seleccione período, nivel, paralelo y materias en el sistema antes de validar el archivo."),
        ("Identificación", "Ingrese como texto, conservando ceros iniciales; no utilice fórmulas."),
        ("Catálogos", "Seleccione el tipo de documento o ingrese su código vigente."),
        ("Obligatorios", "Complete tipo de documento, número de cédula, primer nombre y primer apellido. Segundo nombre y segundo apellido son opcionales."),
        ("Documentos", "Después del ingreso, abra Documentos para cargar el expediente de cada estudiante."),
        ("Office 365 y Moodle", "Los cuatro campos de nombres y apellidos se utilizan directamente para las credenciales. Conserve completos los nombres y apellidos compuestos."),
        ("Datos no informados", "No se inventa un correo personal. Sexo, estado civil y etnia utilizan los valores predeterminados de DATOS_ESTUD; los estudiantes existentes conservan su información."),
    ], start=1):
        _text_cell(parameters, row, 1, key)
        _text_cell(parameters, row, 2, value)
    parameters.column_dimensions["A"].width = 24
    parameters.column_dimensions["B"].width = 115
    references = workbook.create_sheet("Catálogos")
    refs_column = 1
    for index, (field, label, key, required) in enumerate(SIMPLE_COLUMNS, start=1):
        letter = get_column_letter(index)
        cell = _text_cell(students, 1, index, label + (" *" if required else ""))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="176C7C")
        cell.comment = Comment("Obligatorio" if required else "Opcional", "INTEC")
        students.column_dimensions[letter].width = 28 if field not in {"apellidos", "nombres", "correo", "correo_intec"} else 36
        students.column_dimensions[letter].number_format = "@"
        _text_cell(students, 2, index, "")
        if field == "fecha_nacimiento":
            cell.comment = Comment("Fecha: AAAA-MM-DD o DD/MM/AAAA", "INTEC")
        if not key:
            continue
        _text_cell(references, 1, refs_column, label)
        _text_cell(references, 1, refs_column + 1, "Código del padre")
        options = _options(catalogs, key)
        for row, option in enumerate(options, start=2):
            _text_cell(references, row, refs_column, f"{option['value']} - {option['label']}")
            _text_cell(references, row, refs_column + 1, option.get("parent_value", ""))
        references.column_dimensions[get_column_letter(refs_column)].width = 38
        references.column_dimensions[get_column_letter(refs_column + 1)].width = 18
        if options:
            ref_letter = get_column_letter(refs_column)
            range_name = f"catalogo_{field}"
            workbook.defined_names.add(DefinedName(range_name, attr_text=f"'Catálogos'!${ref_letter}$2:${ref_letter}${len(options) + 1}"))
            validation = DataValidation(type="list", formula1=f"={range_name}", allow_blank=not required)
            validation.errorTitle = "Opción no válida"
            validation.error = "Seleccione una opción del catálogo."
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            students.add_data_validation(validation)
            validation.add(f"{letter}2:{letter}1048576")
        refs_column += 2
    students.auto_filter.ref = f"A1:{get_column_letter(len(SIMPLE_COLUMNS))}2"
    pensum = workbook.create_sheet("Pensum")
    pensum.append(["Código de materia", "Código único", "Materia", "Semestre", "Créditos"])
    for row, subject in enumerate(sorted(subjects, key=lambda item: (int(item.get("semestre") or 0), int(item["codigo_materia"]))), start=2):
        for col, value in enumerate([subject["codigo_materia"], subject.get("cod_materia") or "", subject["nombre_materia"], subject.get("semestre") or "", subject.get("creditos") or ""], start=1):
            _text_cell(pensum, row, col, value)
    for letter, width in {"A": 22, "B": 30, "C": 70, "D": 15, "E": 15}.items():
        pensum.column_dimensions[letter].width = width
    references.freeze_panes = pensum.freeze_panes = "A2"
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def _header(value) -> str:
    return " ".join("".join(c for c in unicodedata.normalize("NFKD", str(value or "")) if not unicodedata.combining(c)).replace("*", "").lower().split())


def read_students(content: bytes, career_code: int, catalogs: dict) -> list[dict]:
    if not content or len(content) > MAX_FILE_BYTES:
        raise HTTPException(400, detail="El Excel debe tener contenido y no superar 12 MB.")
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 3000 or sum(entry.file_size for entry in entries) > 64 * 1024 * 1024:
                raise HTTPException(400, detail="El contenido descomprimido del Excel es demasiado grande.")
            if any("vbaproject" in entry.filename.lower() for entry in entries):
                raise HTTPException(400, detail="No se admiten macros en la plantilla.")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, ValueError, KeyError, OSError, ParseError, InvalidFileException) as error:
        raise HTTPException(400, detail="No se pudo leer el archivo. Utilice la plantilla .xlsx de ingreso directo.") from error
    try:
        if "Estudiantes" not in workbook or "Parámetros" not in workbook:
            raise HTTPException(400, detail="Faltan las hojas Estudiantes o Parámetros de la plantilla.")
        parameters = workbook["Parámetros"]
        version = parameters["B1"].value
        if version not in {VERSION, LEGACY_VERSION} or str(parameters["B2"].value) != str(career_code):
            raise HTTPException(400, detail="La plantilla no corresponde a la carrera seleccionada o al formato vigente.")
        students = workbook["Estudiantes"]
        students.reset_dimensions()
        rows = iter(students.iter_rows())
        headers = next(rows, ())
        columns = SIMPLE_COLUMNS if version == VERSION else COLUMNS + CREDENTIAL_COLUMNS
        by_header = {_header(label): (field, label, key, required) for field, label, key, required in columns}
        names = [_header(cell.value) for cell in headers]
        allowed_headers = [set(by_header)]
        if version == LEGACY_VERSION:
            allowed_headers.append({_header(label) for _, label, _, _ in COLUMNS})
        if len(set(names)) != len(names) or set(names) not in allowed_headers:
            raise HTTPException(400, detail="Las columnas no corresponden a la plantilla. Descárguela nuevamente; no agregue ni elimine columnas.")
        parsed = []
        for row_number, cells in enumerate(rows, start=2):
            if not any(cell.value not in (None, "") for cell in cells):
                continue
            data, issues = {}, []
            if any(cell.value not in (None, "") for cell in cells[len(names):]):
                issues.append("La fila contiene columnas adicionales que no pertenecen a la plantilla.")
            for index, name in enumerate(names):
                cell = cells[index] if index < len(cells) else None
                field, label, key, required = by_header[name]
                value = cell.value if cell else None
                if cell and cell.data_type in {"f", "e"}:
                    issues.append(f"{label}: no se admiten fórmulas ni errores de Excel.")
                    value = None
                if field == "identificacion" and value is not None and not isinstance(value, str):
                    issues.append("Identificación: ingrésela como texto para conservar los ceros iniciales.")
                if field == "fecha_nacimiento" and isinstance(value, (date, datetime)):
                    value = value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
                elif isinstance(value, float) and value.is_integer():
                    value = str(int(value))
                value = str(value).strip() if value is not None else ""
                if required and not value:
                    issues.append(f"{label}: es obligatorio.")
                if key and value:
                    options = _options(catalogs, key)
                    option = next((item for item in options if value in {str(item['value']), f"{item['value']} - {item['label']}"}), None)
                    if option:
                        value = str(option["value"])
                    else:
                        issues.append(f"{label}: no pertenece al catálogo vigente.")
                if field == "fecha_nacimiento" and value:
                    try:
                        value = datetime.strptime(value, "%d/%m/%Y").date().isoformat() if "/" in value else date.fromisoformat(value).isoformat()
                    except ValueError:
                        issues.append("Fecha de nacimiento: use AAAA-MM-DD o DD/MM/AAAA.")
                data[field] = value or (None if field in {"fecha_nacimiento", "provincia_residencia"} else "")
            credential_names = {field: data.pop(field, "") for field, _, _, _ in CREDENTIAL_COLUMNS}
            if version == VERSION:
                data.update(
                    nombres=" ".join(f"{credential_names['primer_nombre']} {credential_names['segundo_nombre']}".split()),
                    apellidos=" ".join(f"{credential_names['primer_apellido']} {credential_names['segundo_apellido']}".split()),
                    correo="", correo_intec="",
                )
            parsed.append({"fila": row_number, "datos": data, "errores": issues, "credenciales": credential_names})
        if not parsed:
            raise HTTPException(400, detail="La hoja Estudiantes no contiene registros para ingresar.")
        return parsed
    except (BadZipFile, ValueError, KeyError, OSError, ParseError, InvalidFileException) as error:
        raise HTTPException(400, detail="El archivo contiene datos o estructura de Excel no válidos. Descargue la plantilla nuevamente.") from error
    finally:
        workbook.close()
